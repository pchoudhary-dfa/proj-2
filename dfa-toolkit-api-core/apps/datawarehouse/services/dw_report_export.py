"""
Generic report export engine — CSV, XLSX, PDF, DOCX.
"""
from __future__ import annotations

import csv
import re
from datetime import datetime
from io import BytesIO, StringIO
from typing import Any, Literal, Optional

from apps.datawarehouse.utils.export_utils import (
    build_generic_export_xlsx,
    make_export_filename,
    translate_header,
)

ExportFormat = Literal["csv", "xlsx", "pdf", "docx"]

MEDIA_TYPES: dict[ExportFormat, str] = {
    "csv": "text/csv",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "pdf": "application/pdf",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
}

EXTENSIONS: dict[ExportFormat, str] = {
    "csv": ".csv",
    "xlsx": ".xlsx",
    "pdf": ".pdf",
    "docx": ".docx",
}


def make_report_filename(
    slug: str,
    db_display_name: str,
    fmt: ExportFormat,
) -> str:
    """Generate filename: SLUG_dw_name_YYYY-MM-DD-HH-MM-SS.ext"""
    safe_slug = re.sub(r"[^A-Za-z0-9._-]+", "_", slug.replace("/", "_").upper())
    return make_export_filename(safe_slug, db_display_name, extension=EXTENSIONS[fmt])


def _stringify(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "true" if value else "false"
    if hasattr(value, "isoformat"):
        return value.isoformat()
    return str(value)


def _export_csv(
    column_defs: list[str],
    rows: list[dict[str, Any]],
    language_code: str,
) -> BytesIO:
    headers = [translate_header(col, language_code) for col in column_defs]
    buf = StringIO()
    buf.write("\ufeff")  # UTF-8 BOM for Excel
    writer = csv.writer(buf)
    writer.writerow(headers)
    for row in rows:
        writer.writerow([_stringify(row.get(col)) for col in column_defs])
    output = BytesIO(buf.getvalue().encode("utf-8"))
    output.seek(0)
    return output


def _export_xlsx(
    column_defs: list[str],
    rows: list[dict[str, Any]],
    language_code: str,
    db_display_name: str,
    title: str,
    filename_prefix: str,
    preamble_rows: Optional[list] = None,
) -> tuple[BytesIO, str]:
    data_rows = [[row.get(col) for col in column_defs] for row in rows]
    return build_generic_export_xlsx(
        title, column_defs, data_rows, language_code, db_display_name, filename_prefix,
        preamble_rows=preamble_rows,
    )


def _export_pdf(
    column_defs: list[str],
    rows: list[dict[str, Any]],
    language_code: str,
    title: str,
    db_display_name: str,
) -> BytesIO:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    headers = [translate_header(col, language_code) for col in column_defs]
    table_data = [headers] + [[_stringify(row.get(col)) for col in column_defs] for row in rows]

    page_size = landscape(A4) if len(column_defs) > 4 else A4
    output = BytesIO()
    doc = SimpleDocTemplate(output, pagesize=page_size)
    styles = getSampleStyleSheet()
    ts = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    subtitle = f"{db_display_name} — {ts}"

    elements = [
        Paragraph(title, styles["Heading2"]),
        Paragraph(subtitle, styles["Normal"]),
        Spacer(1, 12),
    ]

    table = Table(table_data, repeatRows=1)
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.grey),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.whitesmoke),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTSIZE", (0, 0), (-1, -1), 8),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                (
                    "ROWBACKGROUNDS",
                    (0, 1),
                    (-1, -1),
                    [colors.white, colors.Color(0.95, 0.95, 0.95)],
                ),
            ]
        )
    )
    elements.append(table)
    doc.build(elements)
    output.seek(0)
    return output


def _export_docx(
    column_defs: list[str],
    rows: list[dict[str, Any]],
    language_code: str,
    title: str,
    db_display_name: str,
) -> BytesIO:
    from docx import Document

    headers = [translate_header(col, language_code) for col in column_defs]
    doc = Document()
    doc.add_heading(title, level=1)
    ts = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    doc.add_paragraph(f"{db_display_name} — {ts}")

    table = doc.add_table(rows=1 + len(rows), cols=len(column_defs))
    table.style = "Table Grid"
    hdr_cells = table.rows[0].cells
    for i, header in enumerate(headers):
        hdr_cells[i].text = header
        for paragraph in hdr_cells[i].paragraphs:
            for run in paragraph.runs:
                run.bold = True

    for row_idx, row in enumerate(rows, start=1):
        cells = table.rows[row_idx].cells
        for col_idx, col in enumerate(column_defs):
            cells[col_idx].text = _stringify(row.get(col))

    output = BytesIO()
    doc.save(output)
    output.seek(0)
    return output


def export_report(
    *,
    column_defs: list[str],
    rows: list[dict[str, Any]],
    fmt: ExportFormat,
    title: str,
    slug: str,
    db_display_name: str,
    language_code: str,
    filename_prefix: Optional[str] = None,
    preamble_rows: Optional[list] = None,
) -> tuple[BytesIO, str, str]:
    """
    Convert report rows to a downloadable file stream.

    Returns (stream, filename, media_type).
    """
    prefix = filename_prefix or slug.replace("/", "_").upper()
    filename = make_report_filename(slug, db_display_name, fmt)

    if fmt == "csv":
        return _export_csv(column_defs, rows, language_code), filename, MEDIA_TYPES["csv"]
    if fmt == "xlsx":
        stream, xlsx_filename = _export_xlsx(
            column_defs, rows, language_code, db_display_name, title, prefix,
            preamble_rows=preamble_rows,
        )
        return stream, xlsx_filename, MEDIA_TYPES["xlsx"]
    if fmt == "pdf":
        return (
            _export_pdf(column_defs, rows, language_code, title, db_display_name),
            filename,
            MEDIA_TYPES["pdf"],
        )
    if fmt == "docx":
        return (
            _export_docx(column_defs, rows, language_code, title, db_display_name),
            filename,
            MEDIA_TYPES["docx"],
        )
    raise ValueError(f"Unsupported export format: {fmt!r}")
