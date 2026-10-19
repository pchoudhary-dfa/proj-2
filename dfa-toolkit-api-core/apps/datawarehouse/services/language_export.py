"""
Language export service.

Builds a single XLSX workbook containing 14 sheets, each comparing a source and target locale
table side-by-side (from language vs to language).
"""

from __future__ import annotations

import logging
import re
from dataclasses import dataclass
from io import BytesIO
from typing import Any, List, Optional, Sequence, Tuple

from sqlalchemy import text
from sqlalchemy.orm import Session

from apps.shared.exceptions import BadRequest
from apps.shared.utils.validators import normalize_language_code_value

logger = logging.getLogger(__name__)
_LANG_CODE_RE = re.compile(r"^[a-z]{2,10}$")


def _sanitize_db_connection_name(name: Optional[str], fallback: str = "dw") -> str:
    """
    Convert a human-readable db name into a safe filename segment.
    Allows: A-Z a-z 0-9 . _ -
    """
    raw = (name or "").strip() or fallback
    raw = raw.replace(" ", "-")
    safe = re.sub(r"[^A-Za-z0-9._-]+", "-", raw).strip("-")
    return safe or fallback


def _normalize_lang_code(code: Optional[str]) -> Optional[str]:
    if code is None:
        return None
    norm = normalize_language_code_value(str(code))
    if not norm:
        return None
    if not _LANG_CODE_RE.match(norm):
        raise BadRequest(f"Invalid language code '{code}'")
    return norm


def _table_exists(dw_db: Session, table_name: str) -> bool:
    row = dw_db.execute(
        text(
            """
            SELECT 1
            FROM information_schema.tables
            WHERE table_schema = DATABASE()
              AND table_name = :t
            LIMIT 1
            """
        ),
        {"t": table_name},
    ).fetchone()
    return bool(row)


def _fetch_language_row(dw_db: Session, lang_code: str) -> Optional[Tuple[str, Optional[str], bool]]:
    """
    Returns (language_code, language_name, language_default) or None.
    """
    row = dw_db.execute(
        text(
            """
            SELECT language_code, language_name, language_default
            FROM ut_language
            WHERE language_code = :code
            LIMIT 1
            """
        ),
        {"code": lang_code},
    ).fetchone()
    if not row:
        return None
    code = str(row[0]) if row[0] is not None else ""
    name = row[1] if len(row) > 1 else None
    is_default = bool(row[2]) if len(row) > 2 else False
    return code, name, is_default


def _resolve_default_language_code(dw_db: Session) -> str:
    raise NotImplementedError("Use _try_resolve_default_language_code")


def _try_resolve_default_language_code(dw_db: Session) -> Optional[str]:
    """
    Try to resolve the DW default language code using ut_language.language_default=1.
    Returns None when no default is configured.

    Note: For language export, when fromLangCode is omitted and no default exists,
    we intentionally skip the from-language column (instead of failing or falling back).
    """
    row = dw_db.execute(
        text(
            """
            SELECT language_code
            FROM ut_language
            WHERE language_default = 1
            LIMIT 1
            """
        )
    ).fetchone()
    if row and row[0]:
        return _normalize_lang_code(str(row[0]))
    return None


@dataclass(frozen=True)
class _SheetSpec:
    sheet_name: str
    table_prefix: str
    pk_col: str
    value_col: str
    id_col: Optional[str] = None  # if None, pk_col is used as identifier
    id_header: str = "GID"  # header label for the identifier column


_SHEETS: Sequence[_SheetSpec] = (
    _SheetSpec(
        sheet_name="INDICATOR_CLASSIFICATIONS",
        table_prefix="ut_indicator_classifications",
        pk_col="ic_nid",
        id_col="ic_gid",
        value_col="ic_name",
    ),
    _SheetSpec(
        sheet_name="SUBGROUP_VALS",
        table_prefix="ut_subgroup_vals",
        pk_col="subgroup_val_nid",
        id_col="subgroup_val_gid",
        value_col="subgroup_val",
    ),
    _SheetSpec(
        sheet_name="INDICATOR",
        table_prefix="ut_indicator",
        pk_col="indicator_nid",
        id_col="indicator_gid",
        value_col="indicator_name",
    ),
    _SheetSpec(
        sheet_name="SUBGROUP",
        table_prefix="ut_subgroup",
        pk_col="subgroup_nid",
        id_col="subgroup_gid",
        value_col="subgroup_name",
    ),
    _SheetSpec(
        sheet_name="UNIT",
        table_prefix="ut_unit",
        pk_col="unit_nid",
        id_col="unit_gid",
        value_col="unit_name",
    ),
    _SheetSpec(
        sheet_name="FOOTNOTE",
        table_prefix="ut_footnote",
        pk_col="footnote_nid",
        id_col="footnote_gid",
        value_col="footnote",
    ),
    _SheetSpec(
        sheet_name="AREA",
        table_prefix="ut_area",
        pk_col="area_nid",
        id_col="area_gid",
        value_col="area_name",
    ),
    _SheetSpec(
        sheet_name="AREA_LEVEL",
        table_prefix="ut_area_level",
        pk_col="level_nid",
        id_col=None,  # PK NID as identifier (per requirement)
        value_col="area_level_name",
        id_header="NID",
    ),
    _SheetSpec(
        sheet_name="AREA_FEATURE_TYPE",
        table_prefix="ut_area_feature_type",
        pk_col="feature_type_nid",
        id_col=None,  # PK NID as identifier (per requirement)
        value_col="feature_type",
        id_header="NID",
    ),
    _SheetSpec(
        sheet_name="MAP_NAME",
        table_prefix="ut_area_map_metadata",
        pk_col="metadata_nid",
        id_col=None,  # PK NID as identifier (per requirement)
        value_col="layer_name",
        id_header="NID",
    ),
    _SheetSpec(
        sheet_name="SUBGROUP_TYPE",
        table_prefix="ut_subgroup_type",
        pk_col="subgroup_type_nid",
        id_col="subgroup_type_gid",
        value_col="subgroup_type_name",
    ),
    _SheetSpec(
        sheet_name="DBMETADATA",
        table_prefix="ut_dbmetadata",
        pk_col="dbmtd_nid",
        id_col=None,  # PK NID as identifier (per requirement)
        value_col="dbmtd_desc",
        id_header="NID",
    ),
    _SheetSpec(
        sheet_name="METADATA_CATEGORY",
        table_prefix="ut_metadata_category",
        pk_col="categorynid",
        id_col="categorygid",
        value_col="categoryname",
    ),
    _SheetSpec(
        sheet_name="METADATA_REPORT",
        table_prefix="ut_metadatareport",
        pk_col="metadatareport_nid",
        id_col=None,  # PK NID as identifier (per requirement)
        value_col="metadata",
        id_header="NID",
    ),
)


def _fetch_sheet_rows(
    dw_db: Session,
    *,
    spec: _SheetSpec,
    from_lang_code: str,
    to_lang_code: str,
) -> List[Tuple[Any, Any, Any]]:
    from_tbl = f"{spec.table_prefix}_{from_lang_code}"
    to_tbl = f"{spec.table_prefix}_{to_lang_code}"

    id_col = spec.id_col or spec.pk_col
    pk_col = spec.pk_col
    val_col = spec.value_col

    # Best-effort include rows present only on one side (MySQL has no FULL OUTER JOIN).
    sql = f"""
        SELECT
            COALESCE(f.`{id_col}`, t.`{id_col}`) AS gid,
            f.`{val_col}` AS from_value,
            t.`{val_col}` AS to_value
        FROM `{from_tbl}` f
        LEFT JOIN `{to_tbl}` t
            ON t.`{pk_col}` = f.`{pk_col}`
        UNION ALL
        SELECT
            COALESCE(f.`{id_col}`, t.`{id_col}`) AS gid,
            f.`{val_col}` AS from_value,
            t.`{val_col}` AS to_value
        FROM `{to_tbl}` t
        LEFT JOIN `{from_tbl}` f
            ON f.`{pk_col}` = t.`{pk_col}`
        WHERE f.`{pk_col}` IS NULL
    """
    rows = dw_db.execute(text(sql)).fetchall()
    out = [(r[0], r[1], r[2]) for r in rows if r]
    out.sort(key=lambda x: ("" if x[0] is None else str(x[0]).lower()))
    return out


def _fetch_sheet_rows_to_only(
    dw_db: Session,
    *,
    spec: _SheetSpec,
    to_lang_code: str,
) -> List[Tuple[Any, Any]]:
    to_tbl = f"{spec.table_prefix}_{to_lang_code}"
    id_col = spec.id_col or spec.pk_col
    val_col = spec.value_col

    sql = f"""
        SELECT
            t.`{id_col}` AS ident,
            t.`{val_col}` AS to_value
        FROM `{to_tbl}` t
    """
    rows = dw_db.execute(text(sql)).fetchall()
    out = [(r[0], r[1]) for r in rows if r]
    out.sort(key=lambda x: ("" if x[0] is None else str(x[0]).lower()))
    return out


def build_languages_export_xlsx(
    *,
    dw_db: Session,
    from_lang_code: Optional[str],
    to_lang_code: str,
    db_connection_display_name: str,
) -> Tuple[BytesIO, str]:
    """
    Build XLSX export for 14 locale tables, comparing from_lang_code and to_lang_code.

    Returns:
        (xlsx_stream, filename)
    """
    to_code = _normalize_lang_code(to_lang_code)
    if not to_code:
        raise BadRequest("toLangCode is required")

    # fromLangCode behavior:
    # - if provided: must exist in ut_language, and we include from-language column
    # - if omitted and DW has a default language: use it and include column
    # - if omitted and no default language exists: skip from-language column
    from_code = _normalize_lang_code(from_lang_code) if from_lang_code is not None else None
    if from_code is None and from_lang_code is None:
        from_code = _try_resolve_default_language_code(dw_db)
    include_from = from_code is not None

    from_lang_name: Optional[str] = None
    if include_from:
        from_row = _fetch_language_row(dw_db, str(from_code))
        if not from_row:
            raise BadRequest(f"Language code '{from_code}' not found in ut_language")
        from_lang_name = (from_row[1] or from_code) if len(from_row) > 1 else str(from_code)

    to_row = _fetch_language_row(dw_db, to_code)
    if not to_row:
        raise BadRequest(f"Language code '{to_code}' not found in ut_language")

    to_lang_name = (to_row[1] or to_code) if len(to_row) > 1 else to_code

    # Validate all required tables exist before building workbook (fail fast).
    missing: List[str] = []
    for spec in _SHEETS:
        tt = f"{spec.table_prefix}_{to_code}"
        if include_from:
            ft = f"{spec.table_prefix}_{from_code}"
            if not _table_exists(dw_db, ft):
                missing.append(ft)
        if not _table_exists(dw_db, tt):
            missing.append(tt)
    if missing:
        shown = ", ".join(missing[:5]) + (" ..." if len(missing) > 5 else "")
        raise BadRequest(f"Required locale tables are missing: {shown}")

    db_connection_name = _sanitize_db_connection_name(db_connection_display_name, fallback="dw")
    if include_from:
        filename = f"{db_connection_name}-Language-{from_lang_name} [{from_code}]-{to_lang_name} [{to_code}].xlsx"
    else:
        filename = f"{db_connection_name}-Language-{to_lang_name} [{to_code}].xlsx"

    from openpyxl import Workbook
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    first = True

    for spec in _SHEETS:
        if first:
            ws = wb.active
            ws.title = spec.sheet_name
            first = False
        else:
            ws = wb.create_sheet(title=spec.sheet_name)

        from_header = f"{from_lang_name} [{from_code}]" if include_from else None
        to_header = f"{to_lang_name} [{to_code}]"
        if include_from:
            headers = [spec.id_header, from_header, to_header]
        else:
            headers = [spec.id_header, to_header]
        ws.append(headers)
        max_lens = [len(str(h)) for h in headers]

        if include_from:
            rows = _fetch_sheet_rows(
                dw_db,
                spec=spec,
                from_lang_code=str(from_code),
                to_lang_code=to_code,
            )
            for ident, from_val, to_val in rows:
                row = [ident, from_val, to_val]
                ws.append(row)
                for idx, val in enumerate(row):
                    sval = "" if val is None else str(val)
                    if len(sval) > max_lens[idx]:
                        max_lens[idx] = len(sval)
        else:
            rows2 = _fetch_sheet_rows_to_only(dw_db, spec=spec, to_lang_code=to_code)
            for ident, to_val in rows2:
                row = [ident, to_val]
                ws.append(row)
                for idx, val in enumerate(row):
                    sval = "" if val is None else str(val)
                    if len(sval) > max_lens[idx]:
                        max_lens[idx] = len(sval)

        padding = 6
        min_width = 12
        max_width = 80
        for i, max_len in enumerate(max_lens, start=1):
            col_letter = get_column_letter(i)
            ws.column_dimensions[col_letter].width = min(
                max(max_len + padding, min_width),
                max_width,
            )

    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output, filename

