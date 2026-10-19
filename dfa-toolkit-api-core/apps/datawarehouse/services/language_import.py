"""
Language import service.

Parses an XLSX file produced by the language export, auto-creates the target
language if it does not exist, and upserts translated values into the
corresponding locale-suffixed tables.
"""

from __future__ import annotations

import logging
import re
from io import BytesIO
from typing import Any, Dict, List, Optional, Tuple

from openpyxl import load_workbook
from sqlalchemy import text
from sqlalchemy.orm import Session

from apps.datawarehouse.db.datawarehouse import (
    credentials_to_schema,
    get_warehouse_credentials,
    get_warehouse_session,
)
from apps.datawarehouse.schemas.language import LanguageCreateSchema
from apps.datawarehouse.services.language_export import (
    _SHEETS,
    _SheetSpec,
    _fetch_language_row,
    _normalize_lang_code,
    _table_exists,
)
from apps.datawarehouse.services.language_service import LanguageService
from apps.shared.exceptions import BadRequest
from apps.shared.utils.validators import normalize_language_code_value

logger = logging.getLogger(__name__)

_HEADER_RE = re.compile(r"^(.+?)\s*\[(\w+)\]$")

_SHEET_NAME_TO_SPEC: Dict[str, _SheetSpec] = {s.sheet_name: s for s in _SHEETS}


def _parse_lang_header(header: str) -> Tuple[str, str]:
    """Parse ``"English [en]"`` into ``("English", "en")``."""
    m = _HEADER_RE.match(header.strip())
    if not m:
        raise BadRequest(
            f"Cannot parse language header '{header}'. "
            "Expected format: 'LanguageName [code]'"
        )
    return m.group(1).strip(), normalize_language_code_value(m.group(2).strip())


def _normalize_cell(value: Any) -> Optional[str]:
    if value is None:
        return None
    s = str(value).strip()
    return s if s else None


def _resolve_pks_by_gid(
    dw_db: Session,
    table_name: str,
    pk_col: str,
    gid_col: str,
    gids: List[str],
) -> Dict[str, int]:
    """Batch-lookup GID -> PK mapping from a locale table."""
    if not gids:
        return {}
    placeholders = ", ".join(f":g{i}" for i in range(len(gids)))
    params = {f"g{i}": g for i, g in enumerate(gids)}
    sql = (
        f"SELECT `{gid_col}`, `{pk_col}` "
        f"FROM `{table_name}` "
        f"WHERE `{gid_col}` IN ({placeholders})"
    )
    rows = dw_db.execute(text(sql), params).fetchall()
    return {str(r[0]): r[1] for r in rows}


def _upsert_row(
    dw_db: Session,
    *,
    from_table: str,
    to_table: str,
    pk_col: str,
    value_col: str,
    pk: Any,
    value: Optional[str],
) -> str:
    """
    Upsert a single row in the target locale table.

    Returns ``"inserted"`` if the row was copied from source and updated,
    ``"updated"`` if an existing row was changed, or ``"unchanged"`` if the
    value was already identical.
    """
    existing = dw_db.execute(
        text(f"SELECT `{value_col}` FROM `{to_table}` WHERE `{pk_col}` = :pk"),
        {"pk": pk},
    ).fetchone()

    if existing is None:
        dw_db.execute(
            text(
                f"INSERT IGNORE INTO `{to_table}` "
                f"SELECT * FROM `{from_table}` WHERE `{pk_col}` = :pk"
            ),
            {"pk": pk},
        )
        dw_db.execute(
            text(
                f"UPDATE `{to_table}` SET `{value_col}` = :val "
                f"WHERE `{pk_col}` = :pk"
            ),
            {"pk": pk, "val": value},
        )
        return "inserted"

    current_value = existing[0]
    if current_value == value:
        return "unchanged"

    dw_db.execute(
        text(
            f"UPDATE `{to_table}` SET `{value_col}` = :val "
            f"WHERE `{pk_col}` = :pk"
        ),
        {"pk": pk, "val": value},
    )
    return "updated"


def import_languages_from_xlsx(
    *,
    file_bytes: bytes,
    dw_id: str,
) -> Dict[str, Any]:
    """
    Import language translations from an XLSX file previously exported by
    ``build_languages_export_xlsx``.

    Returns a summary dict with counts of inserted, updated, and skipped rows.
    """
    if not file_bytes:
        raise BadRequest("Uploaded Excel file is empty")

    wb = load_workbook(filename=BytesIO(file_bytes), data_only=True)

    if not wb.sheetnames:
        raise BadRequest("Workbook contains no sheets")

    # --- Extract language codes from the first sheet's header row -----------
    first_ws = wb[wb.sheetnames[0]]
    header_row = next(first_ws.iter_rows(min_row=1, max_row=1, values_only=True), None)
    if not header_row or len(header_row) < 2:
        raise BadRequest(
            "First sheet must have at least 2 columns "
            "(identifier + target language)"
        )

    num_cols = len(header_row)
    has_from = num_cols >= 3

    if has_from:
        from_lang_name, from_lang_code = _parse_lang_header(str(header_row[1]))
        to_lang_name, to_lang_code = _parse_lang_header(str(header_row[num_cols - 1]))
    else:
        from_lang_code = None
        from_lang_name = None
        to_lang_name, to_lang_code = _parse_lang_header(str(header_row[1]))

    to_lang_code = _normalize_lang_code(to_lang_code)
    if not to_lang_code:
        raise BadRequest("Could not determine target language code from headers")

    if from_lang_code:
        from_lang_code = _normalize_lang_code(from_lang_code)

    # --- Validate from-language and auto-create to-language if needed -------
    cred = get_warehouse_credentials(dw_id)
    warehouse = credentials_to_schema(cred)

    language_created = False

    with get_warehouse_session(warehouse) as dw_db:
        if from_lang_code:
            from_row = _fetch_language_row(dw_db, from_lang_code)
            if not from_row:
                raise BadRequest(
                    f"Source language '{from_lang_code}' does not exist "
                    "in this data warehouse"
                )

        to_row = _fetch_language_row(dw_db, to_lang_code)

    if to_row is None:
        if not from_lang_code:
            raise BadRequest(
                f"Target language '{to_lang_code}' does not exist and no "
                "source language is available to clone from"
            )
        svc = LanguageService()
        display_name = to_lang_name or to_lang_code
        svc.create_language(
            dw_id,
            LanguageCreateSchema(
                language_name=display_name,
                language_code=to_lang_code,
                source_lang=from_lang_code,
                language_default=False,
                language_globallock=False,
            ),
        )
        language_created = True
        logger.info(
            "Auto-created language '%s' in dw_id=%s (cloned from '%s')",
            to_lang_code,
            dw_id,
            from_lang_code,
        )

    # --- Process each sheet ------------------------------------------------
    total_inserted = 0
    total_updated = 0
    total_skipped = 0
    sheets_processed = 0

    with get_warehouse_session(warehouse) as dw_db:
        for sheet_name in wb.sheetnames:
            spec = _SHEET_NAME_TO_SPEC.get(sheet_name)
            if spec is None:
                logger.debug("Skipping unrecognized sheet '%s'", sheet_name)
                continue

            to_table = f"{spec.table_prefix}_{to_lang_code}"
            if not _table_exists(dw_db, to_table):
                logger.warning(
                    "Target table '%s' does not exist, skipping sheet '%s'",
                    to_table,
                    sheet_name,
                )
                continue

            if from_lang_code:
                from_table = f"{spec.table_prefix}_{from_lang_code}"
            else:
                from_table = None

            ws = wb[sheet_name]
            rows = list(ws.iter_rows(min_row=1, values_only=True))
            if len(rows) < 2:
                continue

            ws_header = rows[0]
            data_rows = rows[1:]

            to_col_idx = len(ws_header) - 1

            # Collect identifier -> to_value pairs
            parsed: List[Tuple[Any, Optional[str]]] = []
            for row in data_rows:
                if not row or row[0] is None:
                    total_skipped += 1
                    continue
                identifier = row[0]
                to_value = _normalize_cell(row[to_col_idx]) if to_col_idx < len(row) else None
                parsed.append((identifier, to_value))

            if not parsed:
                continue

            # Resolve identifiers to PKs
            pk_value_pairs: List[Tuple[Any, Optional[str]]] = []

            if spec.id_col is not None:
                # GID-based: batch-resolve GID -> PK via from_lang or to_lang table
                lookup_table = from_table or to_table
                gids = [str(ident) for ident, _ in parsed]
                gid_to_pk = _resolve_pks_by_gid(
                    dw_db, lookup_table, spec.pk_col, spec.id_col, gids
                )
                for ident, to_val in parsed:
                    pk = gid_to_pk.get(str(ident))
                    if pk is None:
                        total_skipped += 1
                        continue
                    pk_value_pairs.append((pk, to_val))
            else:
                # NID-based: identifier IS the PK
                for ident, to_val in parsed:
                    try:
                        pk = int(ident)
                    except (ValueError, TypeError):
                        total_skipped += 1
                        continue
                    pk_value_pairs.append((pk, to_val))

            # Upsert each row
            source_table = from_table or to_table
            for pk, to_val in pk_value_pairs:
                result = _upsert_row(
                    dw_db,
                    from_table=source_table,
                    to_table=to_table,
                    pk_col=spec.pk_col,
                    value_col=spec.value_col,
                    pk=pk,
                    value=to_val,
                )
                if result == "inserted":
                    total_inserted += 1
                elif result == "updated":
                    total_updated += 1

            sheets_processed += 1

    wb.close()

    return {
        "language_created": language_created,
        "to_lang_code": to_lang_code,
        "sheets_processed": sheets_processed,
        "inserted": total_inserted,
        "updated": total_updated,
        "skipped": total_skipped,
    }
