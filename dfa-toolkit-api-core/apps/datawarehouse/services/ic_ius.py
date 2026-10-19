"""
IC–IUS export and import service.

Export: XLSX of leaf IC nodes mapped to IUS (Level path, ClassType, GIDs, dynamic subgroup columns).

Import: XLSX or CSV with the same header layout; upserts ``ut_ic_ius`` and creates missing IC path segments.
"""

from __future__ import annotations

import csv
import logging
import re
from dataclasses import dataclass
from datetime import datetime
from io import BytesIO, StringIO
from typing import Any, Dict, Iterable, List, Optional, Sequence, Tuple

from sqlalchemy.orm import Session

from apps.shared.utils.translations import t_list, t
from apps.datawarehouse.constants import IC_TYPE_ALLOWED_FOR_IC_API
from apps.datawarehouse.db.datawarehouse import get_warehouse_session_from_id
from apps.datawarehouse.models.datawarehouse import (
    IcIus,
    Indicator,
    IndicatorClassification,
    IndicatorUnitSubgroup,
    Unit,
    Subgroup,
    SubgroupCategory,
    SubgroupVal,
    SubgroupValsSubgroup,
    get_default_locale,
)
from apps.datawarehouse.schemas.indicator_classification import IndicatorClassificationCreateSchema
from apps.datawarehouse.services.indicator_classification import IndicatorClassificationService
from apps.datawarehouse.services.ius import resolve_ius_by_indicator_unit_subgroup_gids
from apps.datawarehouse.services.language_service import get_dw_language_codes
from apps.shared.utils.validators import generate_guid
from apps.datawarehouse.utils.export_utils import (
    IC_IUS_IMPORT_MAX_LEVEL_COLUMNS,
    get_ic_ius_import_header_inverse_map,
    normalize_import_header_cell,
)
from apps.shared.utils.log_event import log_background_incident

logger = logging.getLogger(__name__)
_IC_IUS_EXPORT_SOURCE = "apps.datawarehouse.services.ic_ius"


def _sanitize_db_connection_name(name: Optional[str], fallback: str = "dw") -> str:
    """
    Convert a human-readable db name into a safe filename segment.
    Allows: A-Z a-z 0-9 . _ -
    """
    raw = (name or "").strip() or fallback
    raw = raw.replace(" ", "-")
    safe = re.sub(r"[^A-Za-z0-9._-]+", "-", raw).strip("-")
    return safe or fallback


@dataclass(frozen=True)
class _IcRow:
    ic_nid: int
    ic_parent_nid: Optional[int]
    ic_name: str
    ic_type: str


@dataclass(frozen=True)
class _IusRow:
    iusnid: int
    indicator_nid: int
    unit_nid: int
    subgroup_val_nid: int


def _norm_ic_type(v: Any) -> str:
    s = ("" if v is None else str(v)).strip()
    return s


def _is_root_parent(parent_nid: Optional[int]) -> bool:
    return parent_nid in (None, -1, 0)


def _chunked(seq: Sequence[int], size: int) -> Iterable[List[int]]:
    if size <= 0:
        yield list(seq)
        return
    for i in range(0, len(seq), size):
        yield list(seq[i : i + size])


def _safe_str(v: Any) -> str:
    if v is None:
        return ""
    s = str(v)
    return s


def _ic_ius_import_cell_str(val: Any) -> str:
    if val is None:
        return ""
    return str(val).strip()


def _canon_ic_parent_nid(parent_nid: Optional[int]) -> int:
    if parent_nid is None or parent_nid in (-1, 0):
        return 0
    return int(parent_nid)


def _build_ic_ius_import_col_index(
    header_row: Tuple[Any, ...],
    header_inverse: Dict[str, str],
) -> Dict[str, int]:
    col_index: Dict[str, int] = {}
    header_cells = [normalize_import_header_cell(c) for c in header_row]
    for idx, cell in enumerate(header_cells):
        if not cell:
            continue
        field = header_inverse.get(cell)
        if not field:
            continue
        if field in col_index and col_index[field] != idx:
            raise ValueError(f"Duplicate column for field '{field}' in header row")
        col_index[field] = idx
    return col_index


def _max_level_index(col_index: Dict[str, int]) -> int:
    m = 0
    for k in col_index:
        if k.startswith("level_"):
            try:
                m = max(m, int(k.split("_", 1)[1]))
            except (ValueError, IndexError):
                continue
    return m


def _row_dict_from_raw(col_index: Dict[str, int], raw: Tuple[Any, ...]) -> Dict[str, Any]:
    out: Dict[str, Any] = {}
    for field, j in col_index.items():
        if j is None or j >= len(raw):
            out[field] = None
        else:
            out[field] = raw[j]
    return out


def _ic_ius_levels_from_row_dict(row: Dict[str, Any], max_level: int, excel_offset: int) -> List[str]:
    parts: List[str] = []
    for i in range(1, max_level + 1):
        k = f"level_{i}"
        if k not in row:
            break
        parts.append(_ic_ius_import_cell_str(row.get(k)))
    while parts and not parts[-1]:
        parts.pop()
    if not parts:
        return []
    seen_empty = False
    for p in parts:
        if not p:
            seen_empty = True
        elif seen_empty:
            raise ValueError(f"Row {excel_offset}: gap in Level path columns (empty between non-empty levels)")
    return [p for p in parts if p]


def _parse_bool_loose(val: Any) -> Optional[bool]:
    if val is None or val == "":
        return None
    s = str(val).strip().lower()
    if s in ("true", "1", "yes", "y"):
        return True
    if s in ("false", "0", "no", "n"):
        return False
    return None


def _parse_ic_ius_import_xlsx_bytes(
    file_bytes: bytes,
    header_inverse: Dict[str, str],
) -> Tuple[List[Tuple[int, Dict[str, Any]]], int]:
    from openpyxl import load_workbook

    raw_rows: List[Tuple[int, Dict[str, Any]]] = []
    skipped_empty = 0
    col_index: Dict[str, int] = {}

    def _append_data_row(excel_offset: int, raw: Tuple[Any, ...]) -> None:
        nonlocal skipped_empty
        if not raw:
            skipped_empty += 1
            return
        rd = _row_dict_from_raw(col_index, raw)
        ct = _ic_ius_import_cell_str(rd.get("ClassType"))
        ig = _ic_ius_import_cell_str(rd.get("IndicatorGid"))
        ug = _ic_ius_import_cell_str(rd.get("UnitGid"))
        sg = _ic_ius_import_cell_str(rd.get("SubgroupGid"))
        max_lv = _max_level_index(col_index)
        try:
            levels = _ic_ius_levels_from_row_dict(rd, max_lv, excel_offset) if max_lv else []
        except ValueError:
            raise
        if not ct and not ig and not ug and not sg and not levels:
            skipped_empty += 1
            return
        raw_rows.append((excel_offset, rd))

    try:
        workbook = load_workbook(filename=BytesIO(file_bytes), read_only=True, data_only=True)
        try:
            ws = workbook.active
            row_iter = ws.iter_rows(values_only=True)
            header_row = next(row_iter, None)
            if not header_row:
                raise ValueError("Excel file does not contain a header row")
            col_index.update(_build_ic_ius_import_col_index(tuple(header_row), header_inverse))
            required = ("ClassType", "IndicatorGid", "UnitGid", "SubgroupGid", "level_1")
            for r in required:
                if r not in col_index:
                    raise ValueError(
                        f"Missing required column '{r}' (use IC–IUS export layout or translated headers)"
                    )
            excel_off = 2
            for raw in row_iter:
                _append_data_row(excel_off, tuple(raw) if raw else ())
                excel_off += 1
        finally:
            workbook.close()
        return raw_rows, skipped_empty
    except ValueError:
        raise
    except Exception as exc:
        logger.debug("IC-IUS import: read_only XLSX parse failed, retrying standard mode: %s", exc)
        raw_rows.clear()
        skipped_empty = 0
        col_index.clear()

    workbook_fb = load_workbook(filename=BytesIO(file_bytes), data_only=True)
    worksheet = workbook_fb.active
    rows_iter = list(worksheet.iter_rows(values_only=True))
    if not rows_iter:
        raise ValueError("Excel file does not contain a header row")
    col_index.update(_build_ic_ius_import_col_index(tuple(rows_iter[0]), header_inverse))
    required = ("ClassType", "IndicatorGid", "UnitGid", "SubgroupGid", "level_1")
    for r in required:
        if r not in col_index:
            raise ValueError(
                f"Missing required column '{r}' (use IC–IUS export layout or translated headers)"
            )
    excel_off = 2
    for raw in rows_iter[1:]:
        _append_data_row(excel_off, tuple(raw) if raw else ())
        excel_off += 1
    return raw_rows, skipped_empty


def _parse_ic_ius_import_csv_bytes(
    file_bytes: bytes,
    header_inverse: Dict[str, str],
) -> Tuple[List[Tuple[int, Dict[str, Any]]], int]:
    raw_rows: List[Tuple[int, Dict[str, Any]]] = []
    skipped_empty = 0
    text = file_bytes.decode("utf-8-sig", errors="replace")
    reader = csv.reader(StringIO(text))
    rows_list = list(reader)
    if not rows_list:
        raise ValueError("CSV file does not contain a header row")
    col_index = _build_ic_ius_import_col_index(tuple(rows_list[0]), header_inverse)
    required = ("ClassType", "IndicatorGid", "UnitGid", "SubgroupGid", "level_1")
    for r in required:
        if r not in col_index:
            raise ValueError(
                f"Missing required column '{r}' (use IC–IUS export layout or translated headers)"
            )
    max_lv = _max_level_index(col_index)
    for line_no, raw in enumerate(rows_list[1:], start=2):
        if not raw:
            skipped_empty += 1
            continue
        rd = _row_dict_from_raw(col_index, tuple(raw))
        ct = _ic_ius_import_cell_str(rd.get("ClassType"))
        ig = _ic_ius_import_cell_str(rd.get("IndicatorGid"))
        ug = _ic_ius_import_cell_str(rd.get("UnitGid"))
        sg = _ic_ius_import_cell_str(rd.get("SubgroupGid"))
        try:
            levels = _ic_ius_levels_from_row_dict(rd, max_lv, line_no) if max_lv else []
        except ValueError:
            raise
        if not ct and not ig and not ug and not sg and not levels:
            skipped_empty += 1
            continue
        raw_rows.append((line_no, rd))
    return raw_rows, skipped_empty


def detect_ic_ius_import_file_kind(
    file_bytes: bytes,
    filename: Optional[str] = None,
    content_type: Optional[str] = None,
) -> str:
    """Return ``'xlsx'`` or ``'csv'``."""
    ct = (content_type or "").lower()
    if "csv" in ct and "excel" not in ct and "spreadsheet" not in ct:
        return "csv"
    if "spreadsheet" in ct or "excel" in ct or ct.endswith("sheet"):
        return "xlsx"
    if file_bytes[:2] == b"PK":
        return "xlsx"
    fn = (filename or "").lower()
    if fn.endswith(".csv"):
        return "csv"
    if fn.endswith(".xlsx") or fn.endswith(".xlsm"):
        return "xlsx"
    if not file_bytes:
        raise ValueError("Uploaded file is empty")
    try:
        file_bytes[:200].decode("utf-8-sig")
        return "csv"
    except UnicodeDecodeError:
        return "xlsx"


class _IcImportIndex:
    """Mutable IC hierarchy index for one import transaction (non-SR, one locale view)."""

    def __init__(self, ic_by_nid: Dict[int, Any], children: Dict[int, List[int]]) -> None:
        self.ic_by_nid = ic_by_nid
        self.children = children

    @classmethod
    def load(cls, dw_db: Session, language_code: str) -> "_IcImportIndex":
        IcModel = IndicatorClassification(language_code)
        rows = (
            dw_db.query(IcModel)
            .filter((IcModel.ic_type.is_(None)) | (IcModel.ic_type != "SR"))
            .all()
        )
        ic_by_nid: Dict[int, Any] = {}
        children: Dict[int, List[int]] = {}
        for r in rows:
            nid = int(r.ic_nid)
            ic_by_nid[nid] = r
            pkey = _canon_ic_parent_nid(getattr(r, "ic_parent_nid", None))
            children.setdefault(pkey, []).append(nid)
        return cls(ic_by_nid, children)

    def has_non_sr_children(self, ic_nid: int) -> bool:
        kids = self.children.get(ic_nid, [])
        return len(kids) > 0

    def find_under_parent(
        self,
        parent_key: int,
        name: str,
        ic_type_filter: Optional[str],
    ) -> List[int]:
        name_t = name.strip()
        matches: List[int] = []
        for cnid in self.children.get(parent_key, []):
            r = self.ic_by_nid.get(cnid)
            if not r:
                continue
            rn = (_safe_str(getattr(r, "ic_name", None))).strip()
            if rn != name_t:
                continue
            if ic_type_filter is not None:
                rt = (_norm_ic_type(getattr(r, "ic_type", None))).strip().upper()
                if rt != ic_type_filter:
                    continue
            matches.append(cnid)
        return matches

    def register_new_ic(self, ic_nid: int, parent_key: int, orm_row: Any) -> None:
        self.ic_by_nid[ic_nid] = orm_row
        self.children.setdefault(parent_key, []).append(ic_nid)


def import_ic_ius_from_file_bytes(
    dw_id: str,
    file_bytes: bytes,
    *,
    language_code: Optional[str] = None,
    filename: Optional[str] = None,
    content_type: Optional[str] = None,
) -> Dict[str, Any]:
    """
    Import IC–IUS mappings from XLSX or CSV (export-compatible headers).

    Creates missing IC nodes along the Level path (new nodes use the row **ClassType** for ``ic_type``).
    When **IndicatorGid**, **UnitGid**, and **SubgroupGid** are all empty, only the IC hierarchy is
    ensured (no ``ut_ic_ius`` row). When all three GIDs are present, upserts ``ut_ic_ius`` for the
    resolved leaf IC and that IUS (leaf must have no children in the non-SR tree).

    Returns counts: ``ic_mappings_inserted``, ``ic_mappings_updated``, ``ic_nodes_created``, ``skipped_empty_rows``.
    """
    if not file_bytes:
        raise ValueError("Uploaded file is empty")

    header_inverse = get_ic_ius_import_header_inverse_map()
    kind = detect_ic_ius_import_file_kind(file_bytes, filename=filename, content_type=content_type)
    if kind == "csv":
        raw_rows, skipped_empty = _parse_ic_ius_import_csv_bytes(file_bytes, header_inverse)
    else:
        raw_rows, skipped_empty = _parse_ic_ius_import_xlsx_bytes(file_bytes, header_inverse)

    if not raw_rows:
        raise ValueError("No data rows to import")

    ic_svc = IndicatorClassificationService()
    ic_mappings_inserted = 0
    ic_mappings_updated = 0
    ic_nodes_created = 0

    with get_warehouse_session_from_id(dw_id) as dw_db:
        get_dw_language_codes(dw_id)
        lc = language_code if language_code is not None else get_default_locale(dw_db)
        idx = _IcImportIndex.load(dw_db, lc)
        max_lv = IC_IUS_IMPORT_MAX_LEVEL_COLUMNS

        for excel_offset, rd in raw_rows:
            class_type = _ic_ius_import_cell_str(rd.get("ClassType")).upper()
            if not class_type:
                raise ValueError(f"Row {excel_offset}: ClassType is required")
            if class_type not in IC_TYPE_ALLOWED_FOR_IC_API:
                raise ValueError(
                    f"Row {excel_offset}: ClassType {class_type!r} not allowed "
                    f"(must be one of {list(IC_TYPE_ALLOWED_FOR_IC_API)})"
                )
            ig = _ic_ius_import_cell_str(rd.get("IndicatorGid"))
            ug = _ic_ius_import_cell_str(rd.get("UnitGid"))
            sg = _ic_ius_import_cell_str(rd.get("SubgroupGid"))
            has_full_ius = bool(ig and ug and sg)
            has_any_ius = bool(ig or ug or sg)
            if has_any_ius and not has_full_ius:
                raise ValueError(
                    f"Row {excel_offset}: provide all of IndicatorGid, UnitGid, and SubgroupGid "
                    "to create an IUS mapping, or leave all three empty to import the IC path only"
                )

            levels = _ic_ius_levels_from_row_dict(rd, max_lv, excel_offset)
            if not levels:
                raise ValueError(f"Row {excel_offset}: at least Level1 must be set")

            parent_key = 0
            for depth, seg in enumerate(levels):
                is_leaf_seg = depth == len(levels) - 1
                type_f = class_type if is_leaf_seg else None
                matches = idx.find_under_parent(parent_key, seg, type_f)
                if len(matches) > 1:
                    raise ValueError(
                        f"Row {excel_offset}: ambiguous IC name {seg!r} under parent {parent_key}"
                    )
                if len(matches) == 0:
                    create_payload = IndicatorClassificationCreateSchema(
                        ic_name=seg,
                        ic_parent_nid=parent_key,
                        ic_type=class_type,
                        ic_gid=generate_guid(),
                    )
                    new_nid = ic_svc._create_ic_in_session(
                        dw_db, dw_id, create_payload, language_code=lc
                    )
                    dw_db.flush()
                    ic_nodes_created += 1
                    IcModel = IndicatorClassification(lc)
                    new_row = dw_db.query(IcModel).filter(IcModel.ic_nid == new_nid).first()
                    idx.register_new_ic(new_nid, parent_key, new_row)
                    parent_key = new_nid
                else:
                    parent_key = matches[0]

            leaf_nid = parent_key
            if has_full_ius and idx.has_non_sr_children(leaf_nid):
                raise ValueError(
                    f"Row {excel_offset}: target IC (nid={leaf_nid}) is not a leaf; "
                    "IC–IUS mappings are only allowed on leaf classifications"
                )

            if not has_full_ius:
                continue

            ius_row = resolve_ius_by_indicator_unit_subgroup_gids(dw_db, lc, ig, ug, sg)
            if not ius_row:
                raise ValueError(
                    f"Row {excel_offset}: IUS not found for indicator_gid={ig!r}, unit_gid={ug!r}, subgroup_val_gid={sg!r}"
                )
            iusnid = int(ius_row.iusnid)

            opt_order_raw = rd.get("ic_ius_order")
            opt_label_raw = rd.get("ic_ius_label")
            opt_label = _ic_ius_import_cell_str(opt_label_raw) or None

            existing = (
                dw_db.query(IcIus)
                .filter(IcIus.ic_nid == leaf_nid, IcIus.iusnid == iusnid)
                .first()
            )
            if existing is None:
                opt_rec = _parse_bool_loose(rd.get("recommendedsource"))
                order_val = None
                if opt_order_raw not in (None, ""):
                    try:
                        order_val = float(opt_order_raw)
                    except (TypeError, ValueError) as e:
                        raise ValueError(f"Row {excel_offset}: invalid ic_ius_order") from e
                new_map = IcIus(
                    ic_nid=leaf_nid,
                    iusnid=iusnid,
                    recommendedsource=bool(opt_rec) if opt_rec is not None else False,
                    ic_ius_order=order_val,
                    ic_ius_label=opt_label,
                )
                dw_db.add(new_map)
                ic_mappings_inserted += 1
            else:
                did_update = False
                if "ic_ius_order" in rd and opt_order_raw not in (None, ""):
                    try:
                        existing.ic_ius_order = float(opt_order_raw)
                    except (TypeError, ValueError) as e:
                        raise ValueError(f"Row {excel_offset}: invalid ic_ius_order") from e
                    did_update = True
                if "ic_ius_label" in rd:
                    existing.ic_ius_label = opt_label
                    did_update = True
                if "recommendedsource" in rd and rd.get("recommendedsource") not in (None, ""):
                    pr = _parse_bool_loose(rd.get("recommendedsource"))
                    if pr is None:
                        raise ValueError(
                            f"Row {excel_offset}: recommendedsource must be a boolean "
                            "(e.g. true/false, 1/0)"
                        )
                    existing.recommendedsource = bool(pr)
                    did_update = True
                if did_update:
                    ic_mappings_updated += 1

        if ic_mappings_inserted or ic_mappings_updated:
            from apps.datawarehouse.utils.audit_events import buffer_manual_audit

            buffer_manual_audit(
                dw_db,
                action="UPDATE" if ic_mappings_updated else "CREATE",
                module="ic_ius",
                description=(
                    f"IC–IUS import: inserted={ic_mappings_inserted} "
                    f"updated={ic_mappings_updated} ic_nodes_created={ic_nodes_created}"
                ),
                new_value={
                    "ic_mappings_inserted": ic_mappings_inserted,
                    "ic_mappings_updated": ic_mappings_updated,
                    "ic_nodes_created": ic_nodes_created,
                },
            )

    return {
        "ic_mappings_inserted": ic_mappings_inserted,
        "ic_mappings_updated": ic_mappings_updated,
        "ic_nodes_created": ic_nodes_created,
        "skipped_empty_rows": skipped_empty,
    }


def build_ic_ius_export_xlsx(
    dw_db: Session,
    language_code: str,
    db_connection_display_name: str,
) -> Tuple[BytesIO, str]:
    """
    Build the XLSX export for IC–IUS mappings (leaf IC nodes only), excluding IC type 'SR'.

    IMPORTANT: Caller must validate DW exists, and pass an active DW session (`dw_db`).

    Returns:
        (xlsx_stream, filename)
    """
    db_connection_name = _sanitize_db_connection_name(
        db_connection_display_name,
        fallback="dw",
    )
    ts = datetime.now().strftime("%Y-%m-%d-%H-%M-%S")
    filename = f"ICIUS_{db_connection_name}_{ts}.xlsx"

    try:
        # --- Load IC master (exclude SR) ---
        IcModel = IndicatorClassification(language_code)
        ic_rows = (
            dw_db.query(
                IcModel.ic_nid,
                IcModel.ic_parent_nid,
                IcModel.ic_name,
                IcModel.ic_type,
            )
            .filter((IcModel.ic_type.is_(None)) | (IcModel.ic_type != "SR"))
            .all()
        )

        ic_by_nid: Dict[int, _IcRow] = {}
        children_by_parent: Dict[int, List[int]] = {}

        for r in ic_rows:
            ic_nid = int(r[0]) if r[0] is not None else None
            if ic_nid is None:
                continue
            parent_nid = None if r[1] is None else int(r[1])
            ic_type = _norm_ic_type(r[3])
            ic_name = _safe_str(r[2])
            ic_by_nid[ic_nid] = _IcRow(
                ic_nid=ic_nid,
                ic_parent_nid=parent_nid,
                ic_name=ic_name,
                ic_type=ic_type,
            )
            if parent_nid is not None and not _is_root_parent(parent_nid):
                children_by_parent.setdefault(parent_nid, []).append(ic_nid)

        ic_nids = sorted(ic_by_nid.keys())

        # Leaf ICs: no children within the non-SR IC set.
        parent_nids_that_have_children = {int(p) for p in children_by_parent.keys() if p is not None}
        leaf_ic_nids = [nid for nid in ic_nids if nid not in parent_nids_that_have_children]

        # --- IC hierarchy paths (dynamic Level1..LevelN) ---
        path_cache: Dict[int, List[str]] = {}
        visiting: set[int] = set()

        def build_path_names(ic_nid: int) -> List[str]:
            if ic_nid in path_cache:
                return path_cache[ic_nid]
            if ic_nid in visiting:
                # Cycle safety: break recursion; export best-effort.
                return [_safe_str(ic_by_nid.get(ic_nid).ic_name if ic_nid in ic_by_nid else ic_nid)]
            visiting.add(ic_nid)

            row = ic_by_nid.get(ic_nid)
            if row is None:
                visiting.remove(ic_nid)
                return []
            parent_nid = row.ic_parent_nid
            if _is_root_parent(parent_nid) or parent_nid not in ic_by_nid:
                out = [row.ic_name]
            else:
                out = build_path_names(int(parent_nid)) + [row.ic_name]
            visiting.remove(ic_nid)
            path_cache[ic_nid] = out
            return out

        max_depth = 0
        for nid in ic_nids:
            depth = len(build_path_names(nid))
            if depth > max_depth:
                max_depth = depth
        if max_depth <= 0:
            max_depth = 1

        # --- Load leaf IC -> IUS mappings (ut_ic_ius) ---
        if not leaf_ic_nids:
            ic_ius_pairs: List[Tuple[int, int]] = []
        else:
            ic_ius_rows = (
                dw_db.query(IcIus.ic_nid, IcIus.iusnid)
                .filter(IcIus.ic_nid.in_(leaf_ic_nids))
                .filter(IcIus.ic_nid.isnot(None))
                .filter(IcIus.iusnid.isnot(None))
                .all()
            )
            ic_ius_pairs = [(int(r[0]), int(r[1])) for r in ic_ius_rows if r and r[0] is not None and r[1] is not None]

        iusnids = sorted({ius for (_, ius) in ic_ius_pairs})

        # --- Load IUS rows ---
        ius_by_nid: Dict[int, _IusRow] = {}
        if iusnids:
            ius_rows = (
                dw_db.query(
                    IndicatorUnitSubgroup.iusnid,
                    IndicatorUnitSubgroup.indicator_nid,
                    IndicatorUnitSubgroup.unit_nid,
                    IndicatorUnitSubgroup.subgroup_val_nid,
                )
                .filter(IndicatorUnitSubgroup.iusnid.in_(iusnids))
                .all()
            )
            for r in ius_rows:
                if r and r[0] is not None:
                    iusnid = int(r[0])
                    ius_by_nid[iusnid] = _IusRow(
                        iusnid=iusnid,
                        indicator_nid=int(r[1] or 0),
                        unit_nid=int(r[2] or 0),
                        subgroup_val_nid=int(r[3] or 0),
                    )

        indicator_nids = sorted({row.indicator_nid for row in ius_by_nid.values() if row.indicator_nid})
        unit_nids = sorted({row.unit_nid for row in ius_by_nid.values() if row.unit_nid})
        subgroup_val_nids = sorted({row.subgroup_val_nid for row in ius_by_nid.values() if row.subgroup_val_nid})

        # --- Bulk-load masters (locale-aware) ---
        IndicatorModel = Indicator(language_code)
        UnitModel = Unit(language_code)
        SubgroupValModel = SubgroupVal(language_code)
        SubgroupTypeModel = SubgroupCategory(language_code)  # ut_subgroup_type_<locale>
        SubgroupModel = Subgroup(language_code)  # ut_subgroup_<locale>

        indicator_info: Dict[int, Tuple[str, str]] = {}
        if indicator_nids:
            for chunk in _chunked(indicator_nids, 2000):
                rows = (
                    dw_db.query(
                        IndicatorModel.indicator_nid,
                        IndicatorModel.indicator_name,
                        IndicatorModel.indicator_gid,
                    )
                    .filter(IndicatorModel.indicator_nid.in_(chunk))
                    .all()
                )
                for r in rows:
                    nid = int(r[0])
                    indicator_info[nid] = (_safe_str(r[1]), _safe_str(r[2]))

        unit_info: Dict[int, Tuple[str, str]] = {}
        if unit_nids:
            for chunk in _chunked(unit_nids, 2000):
                rows = (
                    dw_db.query(
                        UnitModel.unit_nid,
                        UnitModel.unit_name,
                        UnitModel.unit_gid,
                    )
                    .filter(UnitModel.unit_nid.in_(chunk))
                    .all()
                )
                for r in rows:
                    nid = int(r[0])
                    unit_info[nid] = (_safe_str(r[1]), _safe_str(r[2]))

        subgroup_val_info: Dict[int, Tuple[str, str]] = {}
        if subgroup_val_nids:
            for chunk in _chunked(subgroup_val_nids, 2000):
                rows = (
                    dw_db.query(
                        SubgroupValModel.subgroup_val_nid,
                        SubgroupValModel.subgroup_val,
                        SubgroupValModel.subgroup_val_gid,
                    )
                    .filter(SubgroupValModel.subgroup_val_nid.in_(chunk))
                    .all()
                )
                for r in rows:
                    nid = int(r[0])
                    subgroup_val_info[nid] = (_safe_str(r[1]), _safe_str(r[2]))

        # --- Subgroup type columns (dynamic) ---
        sg_type_rows = (
            dw_db.query(
                SubgroupTypeModel.subgroup_type_nid,
                SubgroupTypeModel.subgroup_type_name,
                SubgroupTypeModel.subgroup_type_order,
            )
            .order_by(
                SubgroupTypeModel.subgroup_type_order.is_(None).asc(),
                SubgroupTypeModel.subgroup_type_order.asc(),
                SubgroupTypeModel.subgroup_type_nid.asc(),
            )
            .all()
        )
        subgroup_types: List[Tuple[int, str]] = []
        for r in sg_type_rows:
            if r and r[0] is not None:
                subgroup_types.append((int(r[0]), _safe_str(r[1]) or f"Type {int(r[0])}"))

        # --- Resolve subgroup dimension values for each subgroup_val_nid ---
        # 1) ut_subgroup_vals_subgroup: subgroup_val_nid -> subgroup_nid(s)
        sv_to_subgroups: Dict[int, List[int]] = {}
        if subgroup_val_nids:
            link_rows = (
                dw_db.query(
                    SubgroupValsSubgroup.subgroup_val_nid,
                    SubgroupValsSubgroup.subgroup_nid,
                )
                .filter(SubgroupValsSubgroup.subgroup_val_nid.in_(subgroup_val_nids))
                .all()
            )
            for r in link_rows:
                if not r or r[0] is None or r[1] is None:
                    continue
                sv = int(r[0])
                sg = int(r[1])
                sv_to_subgroups.setdefault(sv, []).append(sg)

        all_subgroup_nids = sorted({sg for sgs in sv_to_subgroups.values() for sg in sgs})

        # 2) ut_subgroup_<locale>: subgroup_nid -> (subgroup_name, subgroup_type_nid)
        subgroup_info: Dict[int, Tuple[str, int]] = {}
        if all_subgroup_nids:
            for chunk in _chunked(all_subgroup_nids, 2000):
                rows = (
                    dw_db.query(
                        SubgroupModel.subgroup_nid,
                        SubgroupModel.subgroup_name,
                        SubgroupModel.subgroup_type,
                    )
                    .filter(SubgroupModel.subgroup_nid.in_(chunk))
                    .all()
                )
                for r in rows:
                    if r and r[0] is not None:
                        nid = int(r[0])
                        subgroup_info[nid] = (_safe_str(r[1]), int(r[2] or 0))

        # Precompute per subgroup_val_nid: subgroup_type_nid -> string value for the cell
        sv_dim_cells: Dict[int, Dict[int, str]] = {}
        for sv, sgs in sv_to_subgroups.items():
            type_to_names: Dict[int, List[str]] = {}
            for sg_nid in sgs:
                info = subgroup_info.get(sg_nid)
                if not info:
                    continue
                sg_name, sg_type_nid = info
                if not sg_type_nid:
                    continue
                if sg_name:
                    type_to_names.setdefault(sg_type_nid, []).append(sg_name)
            sv_dim_cells[sv] = {
                t: ", ".join(names) for (t, names) in type_to_names.items() if names
            }

        # --- Build XLSX ---
        from openpyxl import Workbook
        from openpyxl.utils import get_column_letter

        wb = Workbook()
        ws = wb.active
        ws.title = "IC_IUS"

        level_prefix = t("Level", locale=language_code)
        level_headers = [f"{level_prefix}{i}" for i in range(1, max_depth + 1)]

        fixed_header_keys = [
            "ClassType",
            "Indicator",
            "IndicatorGid",
            "Unit",
            "UnitGid",
            "Subgroup",
            "SubgroupGid",
        ]
        fixed_headers = t_list(fixed_header_keys, locale=language_code)
        ic_type_header = fixed_headers[0]
        ius_headers = fixed_headers[1:]

        # Subgroup type names are already locale-aware values from ut_subgroup_type_<locale>.
        dim_headers = [name for (_, name) in subgroup_types]

        headers = [ic_type_header, *level_headers, *ius_headers, *dim_headers]
        ws.append(headers)
        max_lens = [len(str(h)) for h in headers]

        # Deterministic ordering helps diff/repro.
        for (ic_nid, iusnid) in sorted(ic_ius_pairs, key=lambda x: (x[0], x[1])):
            ic = ic_by_nid.get(ic_nid)
            ius = ius_by_nid.get(iusnid)
            if ic is None or ius is None:
                continue

            path = build_path_names(ic_nid)
            levels = [(path[i] if i < len(path) else "") for i in range(max_depth)]

            ind_name, ind_gid = indicator_info.get(ius.indicator_nid, ("", ""))
            unit_name, unit_gid = unit_info.get(ius.unit_nid, ("", ""))
            sv_name, sv_gid = subgroup_val_info.get(ius.subgroup_val_nid, ("", ""))

            dim_map = sv_dim_cells.get(ius.subgroup_val_nid, {})
            dim_cells = [dim_map.get(type_nid, "") for (type_nid, _) in subgroup_types]

            row = [
                ic.ic_type,
                *levels,
                ind_name,
                ind_gid,
                unit_name,
                unit_gid,
                sv_name,
                sv_gid,
                *dim_cells,
            ]
            ws.append(row)
            for idx, val in enumerate(row):
                sval = "" if val is None else str(val)
                if len(sval) > max_lens[idx]:
                    max_lens[idx] = len(sval)

        # Auto-fit column widths based on max content length (+padding)
        padding = 6
        min_width = 12
        max_width = 80
        for i, max_len in enumerate(max_lens, start=1):
            col_letter = get_column_letter(i)
            ws.column_dimensions[col_letter].width = min(
                max(max_len + padding, min_width),
                max_width,
            )

        output = BytesIO()
        wb.save(output)
        output.seek(0)
        return output, filename

    except Exception as e:
        log_background_incident(
            message=f"Error exporting IC–IUS XLSX: {e}",
            error_code="IC_IUS_EXPORT_XLSX_FAILED",
            exc=e,
            source=_IC_IUS_EXPORT_SOURCE,
        )
        raise

