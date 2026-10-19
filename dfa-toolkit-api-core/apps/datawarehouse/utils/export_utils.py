"""
Shared utilities for building XLSX exports for master data.
"""
import json
import os
import re
from datetime import datetime
from io import BytesIO
from typing import Any, Dict, Generator, List, Optional, Sequence, Tuple

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

# ---------------------------------------------------------------------------
# Translation helpers
# ---------------------------------------------------------------------------

# utils/ -> datawarehouse/ -> apps/  then shared/translations
_TRANSLATIONS_DIR = os.path.join(os.path.dirname(__file__), "..", "..", "shared", "translations")
_TRANSLATION_CACHE: Dict[str, Dict[str, str]] = {}


def _translations_path(language_code: str) -> str:
    return os.path.join(_TRANSLATIONS_DIR, f"{language_code}.json")


def load_translations(language_code: str) -> Dict[str, str]:
    """Load (and cache) translation dict for a given language code. Falls back to 'en'."""
    if language_code in _TRANSLATION_CACHE:
        return _TRANSLATION_CACHE[language_code]
    path = _translations_path(language_code)
    if os.path.isfile(path):
        with open(path, encoding="utf-8") as f:
            _TRANSLATION_CACHE[language_code] = json.load(f)
    else:
        _TRANSLATION_CACHE[language_code] = {}
    return _TRANSLATION_CACHE[language_code]


def translate_header(key: str, language_code: str) -> str:
    """Return translated column header for the given key and language.
    Falls back to English, then to the raw key if no translation found."""
    trans = load_translations(language_code)
    if key in trans:
        return trans[key]
    en = load_translations("en")
    return en.get(key, key)


def normalize_import_header_cell(value: Any) -> str:
    """Normalize Excel header text for fuzzy matching (strip, collapse spaces, lower ASCII)."""
    if value is None:
        return ""
    s = str(value).strip()
    s = re.sub(r"\s+", " ", s)
    return s.lower()


def list_translation_locale_codes() -> List[str]:
    """Locale codes discovered from filenames `apps/shared/translations/<lc>.json`."""
    if not os.path.isdir(_TRANSLATIONS_DIR):
        return []
    codes: List[str] = []
    for name in sorted(os.listdir(_TRANSLATIONS_DIR)):
        if name.endswith(".json"):
            codes.append(os.path.splitext(name)[0])
    return codes


def build_import_header_inverse_map(column_defs: List[str]) -> Dict[str, str]:
    """
    Map normalized spreadsheet header strings to internal column keys.

    Builds labels from translate_header(key, locale) for every translation JSON locale,
    plus the raw internal key itself (covers English headers).
    Duplicate normalized labels targeting different internal keys raise ValueError.
    """
    inverse: Dict[str, str] = {}
    locales = list_translation_locale_codes()
    sources: List[tuple[str, str]] = []
    for field in column_defs:
        sources.append((normalize_import_header_cell(field), field))
    for lc in locales:
        for field in column_defs:
            sources.append((normalize_import_header_cell(translate_header(field, lc)), field))
    seen_conflicts: Dict[str, str] = {}
    for label, field in sources:
        if not label:
            continue
        if label in inverse and inverse[label] != field:
            raise ValueError(
                f"Ambiguous Excel header normalization: '{label!r}' maps to both "
                f"'{inverse[label]}' and '{field}'"
            )
        inverse[label] = field
    return inverse


_AREA_IMPORT_HEADER_INVERSE_CACHE: Optional[Dict[str, str]] = None


def get_area_import_header_inverse_map() -> Dict[str, str]:
    """
    Cached inverse header map for AREA_COLUMN_DEFS (reused across uploads).
    Stale until process restart if translation JSON files change.
    """
    global _AREA_IMPORT_HEADER_INVERSE_CACHE
    if _AREA_IMPORT_HEADER_INVERSE_CACHE is None:
        _AREA_IMPORT_HEADER_INVERSE_CACHE = build_import_header_inverse_map(AREA_COLUMN_DEFS)
    return _AREA_IMPORT_HEADER_INVERSE_CACHE


# IC–IUS import: fixed keys match export translation keys in apps/shared/translations/*.json
# plus optional ut_ic_ius fields (raw keys; translated aliases come from JSON if present).
IC_IUS_IMPORT_FIXED_FIELD_KEYS = [
    "ClassType",
    "Indicator",
    "IndicatorGid",
    "Unit",
    "UnitGid",
    "Subgroup",
    "SubgroupGid",
    "ic_ius_order",
    "ic_ius_label",
    "recommendedsource",
]

# Max LevelN columns recognized on import (matches practical tree depth).
IC_IUS_IMPORT_MAX_LEVEL_COLUMNS = 32

_IC_IUS_IMPORT_HEADER_INVERSE_CACHE: Optional[Dict[str, str]] = None


def get_ic_ius_import_header_inverse_map() -> Dict[str, str]:
    """
    Inverse map for IC–IUS import: export layout + optional mapping columns + Level1..LevelN.

    Level headers match export: ``translate_header('Level', locale) + str(i)`` per locale.
    Also maps normalized ``level_1`` … ``level_N`` literals for plain-ASCII templates.

    Alternate headers ``ic_ius_is_recommended_source`` and ``recommended_source`` map to ``recommendedsource``.
    """
    global _IC_IUS_IMPORT_HEADER_INVERSE_CACHE
    if _IC_IUS_IMPORT_HEADER_INVERSE_CACHE is not None:
        return _IC_IUS_IMPORT_HEADER_INVERSE_CACHE

    inverse = build_import_header_inverse_map(IC_IUS_IMPORT_FIXED_FIELD_KEYS)

    locales = list_translation_locale_codes()
    for lc in locales:
        level_prefix = translate_header("Level", lc)
        for i in range(1, IC_IUS_IMPORT_MAX_LEVEL_COLUMNS + 1):
            header_label = f"{level_prefix}{i}"
            norm = normalize_import_header_cell(header_label)
            internal = f"level_{i}"
            if norm in inverse and inverse[norm] != internal:
                raise ValueError(
                    f"Ambiguous IC-IUS import header: normalized {norm!r} maps to "
                    f"{inverse[norm]!r} and {internal!r}"
                )
            inverse[norm] = internal

    for i in range(1, IC_IUS_IMPORT_MAX_LEVEL_COLUMNS + 1):
        internal = f"level_{i}"
        norm = normalize_import_header_cell(internal)
        if norm not in inverse:
            inverse[norm] = internal

    for alt_label in ("ic_ius_is_recommended_source", "recommended_source"):
        alt = normalize_import_header_cell(alt_label)
        if alt not in inverse:
            inverse[alt] = "recommendedsource"

    _IC_IUS_IMPORT_HEADER_INVERSE_CACHE = inverse
    return _IC_IUS_IMPORT_HEADER_INVERSE_CACHE


# ---------------------------------------------------------------------------
# Filename / workbook helpers
# ---------------------------------------------------------------------------

def sanitize_db_connection_name(name: str, fallback: str = "dw") -> str:
    """Sanitize connection name for use in filenames. Allows A-Z a-z 0-9 . _ -"""
    raw = (name or "").strip() or fallback
    raw = raw.replace(" ", "-")
    safe = re.sub(r"[^A-Za-z0-9._-]+", "-", raw).strip("-")
    return safe or fallback


def make_export_filename(
    prefix: str,
    db_connection_display_name: str,
    suffix: str = "",
    extension: str = ".xlsx",
) -> str:
    """Generate export filename: PREFIX_sanitized_db_name_YYYY-MM-DD-HH-MM-SS[suffix][ext]."""
    safe_name = sanitize_db_connection_name(db_connection_display_name, fallback="dw")
    ts = datetime.now().strftime("%Y-%m-%d-%H-%M-%S")
    ext = extension if extension.startswith(".") else f".{extension}"
    return f"{prefix}_{safe_name}_{ts}{suffix}{ext}"


def build_simple_xlsx(
    sheet_title: str,
    headers: List[str],
    rows: List[List[Any]],
    min_width: int = 12,
    max_width: int = 80,
    padding: int = 6,
    preamble_rows: Optional[List[List[Any]]] = None,
    text_format_column_indices: Optional[Sequence[int]] = None,
) -> BytesIO:
    """
    Build a single-sheet XLSX from headers and rows.
    Optional ``preamble_rows`` are written above the header row (each row padded/truncated
    to match ``len(headers)``). The header row is bold; freeze panes sit on the first data row.
    Auto-sizes columns, freezes panes, and makes the column header row bold.
    ``text_format_column_indices`` (0-based) forces Excel ``@`` text format on those columns
    for every row so values like ``2019-2020`` are not auto-converted to dates.
    Returns BytesIO (already seeked to 0).
    """
    ncol = len(headers)
    preamble_rows = preamble_rows or []

    def _pad_row(r: List[Any]) -> List[Any]:
        if len(r) >= ncol:
            return list(r[:ncol])
        return list(r) + [""] * (ncol - len(r))

    padded_preamble = [_pad_row(list(r)) for r in preamble_rows]

    wb = Workbook()
    ws = wb.active
    ws.title = sheet_title[:31]  # Excel sheet name max 31 chars

    for prow in padded_preamble:
        ws.append(prow)

    ws.append(headers)
    header_row_idx = len(padded_preamble) + 1

    bold_font = Font(bold=True)
    for col_num in range(1, ncol + 1):
        cell = ws.cell(row=header_row_idx, column=col_num)
        cell.font = bold_font

    first_data_row = header_row_idx + 1
    ws.freeze_panes = ws.cell(row=first_data_row, column=1)

    max_lens = [len(str(h)) for h in headers]
    for prow in padded_preamble:
        for idx, val in enumerate(prow):
            if idx < len(max_lens):
                sval = "" if val is None else str(val)
                if len(sval) > max_lens[idx]:
                    max_lens[idx] = len(sval)
    for row in rows:
        ws.append(row)
        for idx, val in enumerate(row):
            if idx < len(max_lens):
                sval = "" if val is None else str(val)
                if len(sval) > max_lens[idx]:
                    max_lens[idx] = len(sval)
    for i, max_len in enumerate(max_lens, start=1):
        col_letter = get_column_letter(i)
        ws.column_dimensions[col_letter].width = min(
            max(max_len + padding, min_width), max_width
        )
    if text_format_column_indices:
        for r in range(1, ws.max_row + 1):
            for ci in text_format_column_indices:
                if 0 <= ci < ncol:
                    ws.cell(row=r, column=ci + 1).number_format = "@"
    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output


SG_VAL_GRID_EMPTY = "-"


def build_sg_vals_grid_export_xlsx(
    items: List[Dict[str, Any]],
    category_columns: List[Dict[str, Any]],
    language_code: str,
    db_connection_display_name: str,
) -> Tuple[BytesIO, str]:
    """
    Build a pivoted subgroups XLSX matching the FED grid layout:
    first column = subgroup val name; subsequent columns = one per ut_subgroup_type
    (ordered by category display order); cell values = ut_subgroup name or ``-``.
    """
    from apps.datawarehouse.constants import QUERY_INC_SUBGROUPS

    headers = [translate_header("subgroup_val", language_code)]
    headers.extend(
        cat.get("sg_cat_name") or f"Category {cat.get('sg_cat_nid')}"
        for cat in category_columns
    )

    data_rows: List[List[Any]] = []
    for item in items:
        dimension = item.get(QUERY_INC_SUBGROUPS) or {}
        by_cat_nid: Dict[int, str] = {}
        for val_data in dimension.values():
            cat_nid = val_data.get("sg_cat_nid")
            if cat_nid is None:
                continue
            sg_val_name = val_data.get("sg_val_name")
            if sg_val_name is not None and str(sg_val_name).strip():
                by_cat_nid[int(cat_nid)] = str(sg_val_name).strip()

        row: List[Any] = [item.get("subgroup_val") or ""]
        for cat in category_columns:
            cat_nid = int(cat["sg_cat_nid"])
            row.append(by_cat_nid.get(cat_nid, SG_VAL_GRID_EMPTY))
        data_rows.append(row)

    output = build_simple_xlsx("Subgroups", headers, data_rows)
    filename = make_export_filename("SUBGROUPS", db_connection_display_name)
    return output, filename


def build_generic_export_xlsx(
    sheet_title: str,
    column_defs: List[str],
    rows: List[List[Any]],
    language_code: str,
    db_connection_display_name: str,
    filename_prefix: str,
    preamble_rows: Optional[List[List[Any]]] = None,
) -> Tuple[BytesIO, str]:
    """
    Build a translated single-sheet XLSX from column_defs and row data.

    column_defs: list of internal field names — each is used directly as the translation key.
    rows: list of lists, each inner list aligned to column_defs order.
    preamble_rows: optional rows above the translated header (same width as column_defs; shorter
        rows are padded with empty strings).
    Returns (xlsx_stream, filename).
    """
    headers = [translate_header(col, language_code) for col in column_defs]
    output = build_simple_xlsx(sheet_title, headers, rows, preamble_rows=preamble_rows)
    filename = make_export_filename(filename_prefix, db_connection_display_name)
    return output, filename


# ---------------------------------------------------------------------------
# Per-entity column definitions — each entry is the internal field name,
# which doubles as the translation key in the JSON files.
# ---------------------------------------------------------------------------

INDICATOR_COLUMN_DEFS = [
    "indicator_nid",
    "indicator_gid",
    "indicator_name",
    "indicator_order",
    "indicator_global",
    "data_exist",
    "highisgood",
    "short_name",
    "keywords",
    # IUS-related (locale fields from related models)
    "unit_nid",
    "unit_gid",
    "unit_name",
    "subgroup_val_nid",
    "subgroup_val_gid",
    "subgroup_val",
    # IC-related (locale fields from related IC model)
    "ic_names",
    "ic_types",
]

UNIT_COLUMN_DEFS = [
    "unit_nid",
    "unit_gid",
    "unit_name",
    "unit_global",
    "unit_format",
]

SG_VAL_COLUMN_DEFS = [
    "subgroup_val_nid",
    "subgroup_val_gid",
    "subgroup_val",
    "subgroup_val_global",
    "subgroup_val_order",
    "is_system",
    # Composition dimension (locale fields from SubgroupCategory + Subgroup models)
    "sg_cat_nid",
    "sg_cat_name",
    "sg_cat_gid",
    "sg_val_nid",
    "sg_val_name",
    "sg_val_gid",
]

INDICATOR_CLASSIFICATION_COLUMN_DEFS = [
    "ic_nid",
    "ic_gid",
    "ic_name",
    "ic_type",
    "ic_parent_nid",
    "ic_short_name",
    "ic_order",
    "ic_global",
]

AREA_COLUMN_DEFS = [
    "area_nid",
    "area_id",
    "area_name",
    "area_gid",
    "area_level",
    "area_parent_nid",
    "area_is_global",
    "area_has_data",
    # Related locale field from AreaLevel model
    "area_level_name",
]

METADATA_CATEGORY_COLUMN_DEFS = [
    "categorynid",
    "categorygid",
    "categoryname",
    "categorytype",
    "categoryorder",
    "parentcategorynid",
    "categorydescription",
    "ispresentational",
    "ismandatory",
]

DATA_POINT_COLUMN_DEFS = [
    "dp_indicator",       # → i_name
    "dp_subgroup",        # → sg_name
    "dp_area",            # → ar_name
    "dp_timeperiod",      # → tp_timeperiod
    "dp_data_value",      # → data_value / textual_data_value
    "dp_sources",         # → src_name
    "dp_footnote",        # → footnote
    "dp_status",          # → review_status
    "dp_data_point_id",   # → data_point_id (used for import round-trips)
]

DATA_COLUMN_DEFS = [
    "data_nid",
    "iusnid",
    "ius_gid",
    "timeperiod_nid",
    "tp_timeperiod",
    "area_nid",
    "ar_id",
    "ar_name",
    "indicator_nid",
    "i_gid",
    "i_name",
    "unit_nid",
    "u_gid",
    "u_name",
    "subgroup_val_nid",
    "sg_gid",
    "sg_name",
    "data_value",
    "textual_data_value",
]

LOG_SERVER_ERROR_COLUMN_DEFS = [
    "lse_id",
    "lse_user_id",
    "lse_dw_id",
    "lse_method",
    "lse_path",
    "lse_status_code",
    "lse_level",
    "lse_error_message",
    "lse_stack_trace",
    "lse_request_body",
    "lse_request_headers",
    "lse_created",
]

LOG_TRANSACTION_COLUMN_DEFS = [
    "lt_id",
    "lt_path",
    "lt_module",
    "lt_identifier",
    "lt_action",
    "lt_status",
    "lt_user_id",
    "lt_dw_id",
    "lt_app_key",
    "lt_dashboard_id",
    "lt_description",
    "lt_previous_value",
    "lt_new_value",
    "lt_created",
]


# ---------------------------------------------------------------------------
# Parallel XLSX export streaming (single file or ZIP)
# ---------------------------------------------------------------------------

_READ_CHUNK = 65_536


def _cancel_pending_export_futures(futures: dict) -> None:
    """Cancel pool jobs not yet started (running workers finish on their own)."""
    for fut in futures:
        if not fut.done():
            fut.cancel()


def stream_single_xlsx_file(worker_args: Any) -> Generator[bytes, None, None]:
    """
    Generate one XLSX via the shared pool and stream file bytes to the client.

    Used when the export fits in a single part (no ZIP wrapper).
    On client disconnect, pending work is cancelled and the temp file removed.
    """
    import os
    from concurrent.futures import Future

    from apps.shared.utils.export_pool import get_export_pool
    from apps.datawarehouse.utils.xlsx_export_worker import worker_generate_xlsx_chunk

    pool = get_export_pool()
    fut: Future[str] = pool.submit(worker_generate_xlsx_chunk, worker_args)
    futures = {fut: worker_args}
    tmp_path: Optional[str] = None

    try:
        tmp_path = fut.result()
        with open(tmp_path, "rb") as fh:
            while True:
                chunk = fh.read(_READ_CHUNK)
                if not chunk:
                    break
                yield chunk
    except GeneratorExit:
        _cancel_pending_export_futures(futures)
        raise
    finally:
        _cancel_pending_export_futures(futures)
        if tmp_path:
            try:
                os.unlink(tmp_path)
            except OSError:
                pass


def stream_parallel_xlsx_zip(
    worker_args_list: List[Any],
    zip_filename: str,
) -> Generator[bytes, None, None]:
    """
    Dispatch worker_args_list to the shared ProcessPoolExecutor, collect
    completed XLSX temp files with as_completed(), and stream them into a
    ZIP via zipstream-ng.

    Uses ZIP_STORED compression (no CPU overhead; XLSX files are already
    compressed internally).

    File order in the ZIP is non-deterministic (whichever worker finishes
    first is added first) — this is intentional per requirements.

    Cleans up temp files after each is streamed into the ZIP.
    On client disconnect, cancels futures that have not started yet.
    Raises any worker exception immediately so the HTTP layer can log it.
    """
    import os
    from concurrent.futures import as_completed
    from zipstream import ZipStream

    from apps.shared.utils.export_pool import get_export_pool
    from apps.datawarehouse.utils.xlsx_export_worker import worker_generate_xlsx_chunk

    pool = get_export_pool()
    futures = {
        pool.submit(worker_generate_xlsx_chunk, args): args
        for args in worker_args_list
    }

    try:
        zs = ZipStream(compress_type=0)

        for fut in as_completed(futures):
            tmp_path: str = fut.result()
            try:
                fname = os.path.basename(tmp_path)
                zs.add_path(tmp_path, fname)
                yield from zs.all_files()
            finally:
                try:
                    os.unlink(tmp_path)
                except OSError:
                    pass

        yield from zs.footer()
    except GeneratorExit:
        _cancel_pending_export_futures(futures)
        raise
    finally:
        _cancel_pending_export_futures(futures)
