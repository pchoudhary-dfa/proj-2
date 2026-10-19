"""
Batched import of data-point values from XLSX/CSV export layout.

Sync: stream rows, bulk_update_mappings per batch with per-batch commit.
Async: queue N system_bg_job_queue rows (shared request_id, dataset_id); any pod can process.
"""
from __future__ import annotations

import csv
import json
import logging
import math
import uuid
from collections import defaultdict
from dataclasses import dataclass
from datetime import datetime, timezone
from io import BytesIO, StringIO
from typing import Any, Generator, Iterable, Literal, Optional

from openpyxl import load_workbook
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from apps.app_db.models.core import Dataset, DatasetDataPoint, SystemBgJobQueue
from apps.shared.config.config import settings
from apps.shared.services.bg_job_dispatch import (
    JOB_TYPE_IMPORT_DATASET_DATA_POINTS_BATCH,
    dispatch_bg_job,
)

logger = logging.getLogger(__name__)

IMPORT_COL_DATA_VALUE = 4
IMPORT_COL_FOOTNOTE = 6
IMPORT_COL_DATA_POINT_ID = 8
IMPORT_EXPECTED_COLS = 9

IMPORT_INVALID_FORMAT = "INVALID_IMPORT_FORMAT"
_EDITABLE_STATUSES = frozenset({"DRAFT", "REVISION"})
_MAX_ERRORS = 100
_IMPORT_BATCH_SIZE = 7000
_TERMINAL_JOB_STATUSES = frozenset({"completed", "failed"})
_ACTIVE_JOB_STATUSES = frozenset({"pending", "processing"})

ImportFormat = Literal["xlsx", "csv"]


@dataclass(frozen=True)
class _ParsedImportRow:
    data_point_id: int
    raw_value: Any
    footnote_raw: Any


def _cell_str(cell_value: Any) -> Optional[str]:
    if cell_value is None:
        return None
    s = str(cell_value).strip()
    return s if s else None


def parse_import_data_point_id(raw: Any) -> Optional[int]:
    if raw is None:
        return None
    if isinstance(raw, bool):
        return None
    if isinstance(raw, int):
        return raw
    if isinstance(raw, float):
        if not math.isfinite(raw) or not raw.is_integer():
            return None
        return int(raw)
    s = str(raw).strip()
    if not s:
        return None
    try:
        return int(s, 10)
    except ValueError:
        pass
    try:
        f = float(s)
        if not math.isfinite(f) or not f.is_integer():
            return None
        return int(f)
    except ValueError:
        return None


def assert_import_allowed(
    db: Session,
    *,
    dataset_id: int,
    dw_id: int,
    check_readiness: bool = True,
    check_editable: bool = True,
) -> Dataset:
    d = (
        db.query(Dataset)
        .filter(Dataset.dataset_id == dataset_id, Dataset.dw_id == dw_id)
        .first()
    )
    if not d:
        raise ValueError("DATASET_NOT_FOUND")
    if check_editable and d.dataset_status not in _EDITABLE_STATUSES:
        raise ValueError("DATASET_NOT_EDITABLE")
    if check_readiness and d.datapoint_readiness in ("CREATED", "UPDATED"):
        raise ValueError("DATAPOINT_POPULATION_IN_PROGRESS")
    return d


def should_import_async(estimated_rows: int) -> bool:
    threshold = settings.DATASET_IMPORT_BG_THRESHOLD
    return threshold > 0 and estimated_rows > threshold


def load_locked_ids(db: Session, dataset_id: int) -> set[int]:
    rows = db.execute(
        select(DatasetDataPoint.data_point_id).where(
            DatasetDataPoint.dataset_id == dataset_id,
            DatasetDataPoint.is_locked.is_(True),
        )
    ).scalars()
    return set(rows)


def _iter_xlsx_rows(file_bytes: bytes) -> Generator[tuple[Any, ...], None, None]:
    wb = None
    try:
        try:
            wb = load_workbook(filename=BytesIO(file_bytes), read_only=True, data_only=True)
        except Exception:
            raise ValueError(IMPORT_INVALID_FORMAT)
        ws = wb.active
        if ws is None:
            raise ValueError(IMPORT_INVALID_FORMAT)
        yield from ws.iter_rows(values_only=True)
    finally:
        if wb is not None:
            wb.close()


def _iter_csv_rows(file_bytes: bytes) -> Generator[list[str], None, None]:
    try:
        text = file_bytes.decode("utf-8-sig")
    except UnicodeDecodeError:
        raise ValueError(IMPORT_INVALID_FORMAT)
    yield from csv.reader(StringIO(text))


def iter_import_rows(
    file_bytes: bytes,
    fmt: ImportFormat,
) -> Generator[tuple[Any, ...], None, None]:
    if fmt == "xlsx":
        yield from _iter_xlsx_rows(file_bytes)
    else:
        yield from _iter_csv_rows(file_bytes)


def _parse_row_tuple(row: tuple[Any, ...]) -> Optional[_ParsedImportRow]:
    if len(row) < IMPORT_EXPECTED_COLS:
        return None
    dp_id = parse_import_data_point_id(row[IMPORT_COL_DATA_POINT_ID])
    if dp_id is None:
        return None
    return _ParsedImportRow(
        data_point_id=dp_id,
        raw_value=row[IMPORT_COL_DATA_VALUE],
        footnote_raw=row[IMPORT_COL_FOOTNOTE],
    )


def count_import_data_rows(file_bytes: bytes, fmt: ImportFormat) -> int:
    count = 0
    for row in iter_import_rows(file_bytes, fmt):
        if _parse_row_tuple(row) is not None:
            count += 1
    return count


def _value_fields_from_cell(raw_value: Any) -> dict[str, Any]:
    from apps.datawarehouse.services.data_service import _process_data_value

    if raw_value is None or str(raw_value).strip() == "":
        return _process_data_value(None)
    return _process_data_value(raw_value)


def _row_to_payload_dict(parsed: _ParsedImportRow) -> dict[str, Any]:
    vf = _value_fields_from_cell(parsed.raw_value)
    return {
        "data_point_id": parsed.data_point_id,
        "data_value": vf["data_value"],
        "textual_data_value": vf["textual_data_value"],
        "istextualdata": vf["istextualdata"],
        "footnote_text": _cell_str(parsed.footnote_raw),
    }


def build_import_batches(
    file_bytes: bytes,
    fmt: ImportFormat,
) -> list[list[dict[str, Any]]]:
    """Parse file and split into payload row batches (normalized value fields)."""
    batches: list[list[dict[str, Any]]] = []
    current: list[dict[str, Any]] = []
    for row in iter_import_rows(file_bytes, fmt):
        parsed = _parse_row_tuple(row)
        if parsed is None:
            continue
        current.append(_row_to_payload_dict(parsed))
        if len(current) >= _IMPORT_BATCH_SIZE:
            batches.append(current)
            current = []
    if current:
        batches.append(current)
    return batches


def _empty_summary(*, import_async: bool = False, **extra: Any) -> dict[str, Any]:
    out: dict[str, Any] = {
        "total_rows": 0,
        "updated": 0,
        "skipped_locked": 0,
        "skipped_no_id": 0,
        "errors": [],
        "import_async": import_async,
        "estimated_rows": extra.pop("estimated_rows", None),
        "request_id": extra.pop("request_id", None),
        "total_batches": extra.pop("total_batches", None),
        "datapoint_readiness": extra.pop("datapoint_readiness", None),
    }
    out.update(extra)
    return out


def _run_batched_import(
    db: Session,
    *,
    dataset_id: int,
    user_id: int,
    rows: Iterable[tuple[Any, ...]],
    max_errors: int = _MAX_ERRORS,
) -> dict[str, Any]:
    locked_ids = load_locked_ids(db, dataset_id)
    now = datetime.now(timezone.utc)

    total_rows = 0
    updated = 0
    skipped_locked = 0
    skipped_no_id = 0
    errors: list[dict] = []
    pending: list[_ParsedImportRow] = []

    def _flush_batch() -> None:
        nonlocal updated, skipped_locked
        if not pending:
            return

        batch_ids = [r.data_point_id for r in pending]
        valid_ids = set(
            db.execute(
                select(DatasetDataPoint.data_point_id).where(
                    DatasetDataPoint.dataset_id == dataset_id,
                    DatasetDataPoint.data_point_id.in_(batch_ids),
                    DatasetDataPoint.is_deleted.is_(False),
                )
            ).scalars()
        )

        mappings: list[dict[str, Any]] = []
        for parsed in pending:
            dp_id = parsed.data_point_id
            if dp_id in locked_ids:
                skipped_locked += 1
                continue
            if dp_id not in valid_ids:
                if len(errors) < max_errors:
                    errors.append({
                        "data_point_id": dp_id,
                        "error": "NOT_FOUND_IN_DATASET",
                        "detail": (
                            f"data_point_id {dp_id} does not belong to dataset {dataset_id}"
                        ),
                    })
                continue

            value_fields = _value_fields_from_cell(parsed.raw_value)
            mappings.append({
                "data_point_id": dp_id,
                "data_value": value_fields["data_value"],
                "textual_data_value": value_fields["textual_data_value"],
                "istextualdata": value_fields["istextualdata"],
                "footnote_text": _cell_str(parsed.footnote_raw),
                "modified": now,
                "modified_user_id": user_id,
            })

        pending.clear()
        if mappings:
            db.bulk_update_mappings(DatasetDataPoint, mappings)
            db.commit()
            updated += len(mappings)

    for row in rows:
        if len(row) < IMPORT_EXPECTED_COLS:
            continue
        parsed = _parse_row_tuple(row)
        if parsed is None:
            skipped_no_id += 1
            continue
        total_rows += 1
        pending.append(parsed)
        if len(pending) >= _IMPORT_BATCH_SIZE:
            _flush_batch()

    _flush_batch()

    return _empty_summary(
        total_rows=total_rows,
        updated=updated,
        skipped_locked=skipped_locked,
        skipped_no_id=skipped_no_id,
        errors=errors,
    )


def import_data_points_sync(
    db: Session,
    *,
    dw_id: int,
    dataset_id: int,
    file_bytes: bytes,
    user_id: int,
    fmt: ImportFormat,
) -> dict[str, Any]:
    assert_import_allowed(db, dataset_id=dataset_id, dw_id=dw_id)
    return _run_batched_import(
        db,
        dataset_id=dataset_id,
        user_id=user_id,
        rows=iter_import_rows(file_bytes, fmt),
    )


def queue_async_import(
    db: Session,
    *,
    dw_id: int,
    dataset_id: int,
    user_id: int,
    batches: list[list[dict[str, Any]]],
    estimated_rows: int,
) -> dict[str, Any]:
    """
    Queue one background job per batch. Sets datapoint_readiness to UPDATED.
    Rolls back readiness if job insert fails.
    """
    if not batches:
        raise ValueError("IMPORT_EMPTY")

    dataset = assert_import_allowed(db, dataset_id=dataset_id, dw_id=dw_id)
    request_id = str(uuid.uuid4())
    total_batches = len(batches)
    prior_readiness = dataset.datapoint_readiness
    dataset.datapoint_readiness = "UPDATED"
    db.flush()

    job_ids: list[int] = []
    try:
        for batch_index, rows in enumerate(batches):
            payload = {
                "request_id": request_id,
                "dataset_id": dataset_id,
                "dw_id": dw_id,
                "user_id": user_id,
                "batch_index": batch_index,
                "total_batches": total_batches,
                "rows": rows,
            }
            result = dispatch_bg_job(
                db=db,
                job_type=JOB_TYPE_IMPORT_DATASET_DATA_POINTS_BATCH,
                payload=payload,
                dw_id=str(dw_id),
                priority=1,
                request_id=request_id,
                dataset_id=dataset_id,
                max_attempts=1,
                commit=False,
            )
            if result.db_job_id is not None:
                job_ids.append(result.db_job_id)
        db.commit()
    except Exception:
        db.rollback()
        dataset = (
            db.query(Dataset)
            .filter(Dataset.dataset_id == dataset_id, Dataset.dw_id == dw_id)
            .first()
        )
        if dataset:
            dataset.datapoint_readiness = prior_readiness
            db.commit()
        raise

    return {
        "import_async": True,
        "request_id": request_id,
        "total_batches": total_batches,
        "estimated_rows": estimated_rows,
        "job_ids": job_ids,
        "datapoint_readiness": "UPDATED",
        "total_rows": 0,
        "updated": 0,
        "skipped_locked": 0,
        "skipped_no_id": 0,
        "errors": [],
    }


def run_import_batch(db: Session, *, payload: dict) -> dict[str, Any]:
    """
    Apply one batch atomically. Returns success result dict or failure dict with failed=True.
    """
    dataset_id = int(payload["dataset_id"])
    user_id = int(payload["user_id"])
    batch_index = int(payload.get("batch_index", 0))
    total_batches = int(payload.get("total_batches", 1))
    rows: list[dict[str, Any]] = list(payload.get("rows") or [])

    locked_ids = load_locked_ids(db, dataset_id)
    now = datetime.now(timezone.utc)

    if not rows:
        return {
            "batch_index": batch_index,
            "total_batches": total_batches,
            "updated": 0,
            "skipped_locked": 0,
            "errors": [],
        }

    batch_ids = [int(r["data_point_id"]) for r in rows]
    valid_ids = set(
        db.execute(
            select(DatasetDataPoint.data_point_id).where(
                DatasetDataPoint.dataset_id == dataset_id,
                DatasetDataPoint.data_point_id.in_(batch_ids),
                DatasetDataPoint.is_deleted.is_(False),
            )
        ).scalars()
    )

    mappings: list[dict[str, Any]] = []
    skipped_locked = 0

    try:
        for row in rows:
            dp_id = int(row["data_point_id"])
            if dp_id in locked_ids:
                skipped_locked += 1
                continue
            if dp_id not in valid_ids:
                err = {
                    "data_point_id": dp_id,
                    "error": "NOT_FOUND_IN_DATASET",
                    "detail": (
                        f"data_point_id {dp_id} does not belong to dataset {dataset_id}"
                    ),
                }
                db.rollback()
                return {
                    "failed": True,
                    "batch_index": batch_index,
                    "total_batches": total_batches,
                    "error_code": "BATCH_IMPORT_FAILED",
                    "message": err["detail"],
                    "errors": [err],
                    "updated": 0,
                    "skipped_locked": skipped_locked,
                }

            mappings.append({
                "data_point_id": dp_id,
                "data_value": row.get("data_value"),
                "textual_data_value": row.get("textual_data_value"),
                "istextualdata": row.get("istextualdata"),
                "footnote_text": row.get("footnote_text"),
                "modified": now,
                "modified_user_id": user_id,
            })

        if mappings:
            db.bulk_update_mappings(DatasetDataPoint, mappings)
        db.commit()
        return {
            "batch_index": batch_index,
            "total_batches": total_batches,
            "updated": len(mappings),
            "skipped_locked": skipped_locked,
            "errors": [],
        }
    except Exception as exc:
        db.rollback()
        return {
            "failed": True,
            "batch_index": batch_index,
            "total_batches": total_batches,
            "error_code": "BATCH_IMPORT_FAILED",
            "message": str(exc)[:2000],
            "errors": [],
            "updated": 0,
            "skipped_locked": skipped_locked,
        }


def maybe_finalize_import_request(
    db: Session,
    *,
    request_id: str,
    dataset_id: int,
    dw_id: int,
    exclude_job_id: Optional[int] = None,
) -> None:
    """Set dataset.datapoint_readiness when all batch jobs for request_id are terminal."""
    active_q = db.query(func.count(SystemBgJobQueue.id)).filter(
        SystemBgJobQueue.request_id == request_id,
        SystemBgJobQueue.dataset_id == dataset_id,
        SystemBgJobQueue.job_type == JOB_TYPE_IMPORT_DATASET_DATA_POINTS_BATCH,
        SystemBgJobQueue.job_status.in_(_ACTIVE_JOB_STATUSES),
    )
    if exclude_job_id is not None:
        active_q = active_q.filter(SystemBgJobQueue.id != exclude_job_id)
    active = active_q.scalar()
    if active and int(active) > 0:
        logger.warning(
            "import finalize deferred request_id=%s dataset_id=%s: %s active batch job(s)",
            request_id,
            dataset_id,
            int(active),
        )
        return

    failed = (
        db.query(func.count(SystemBgJobQueue.id))
        .filter(
            SystemBgJobQueue.request_id == request_id,
            SystemBgJobQueue.dataset_id == dataset_id,
            SystemBgJobQueue.job_type == JOB_TYPE_IMPORT_DATASET_DATA_POINTS_BATCH,
            SystemBgJobQueue.job_status == "failed",
        )
        .scalar()
    ) or 0

    dataset = (
        db.query(Dataset)
        .filter(Dataset.dataset_id == dataset_id, Dataset.dw_id == dw_id)
        .with_for_update()
        .first()
    )
    if not dataset:
        dataset = (
            db.query(Dataset)
            .filter(Dataset.dataset_id == dataset_id)
            .with_for_update()
            .first()
        )
        if dataset is not None and int(dataset.dw_id) != int(dw_id):
            logger.warning(
                "import finalize: dataset_id=%s dw_id=%s does not match dataset.dw_id=%s; "
                "using dataset row",
                dataset_id,
                dw_id,
                dataset.dw_id,
            )
    if not dataset:
        logger.warning(
            "import finalize skipped request_id=%s: dataset not found "
            "(dataset_id=%s dw_id=%s)",
            request_id,
            dataset_id,
            dw_id,
        )
        return

    if int(failed) > 0:
        dataset.datapoint_readiness = "FAILED"
    else:
        dataset.datapoint_readiness = "COMPLETED"


def _parse_json_field(raw: Optional[str]) -> Optional[dict[str, Any]]:
    if not raw:
        return None
    try:
        val = json.loads(raw)
        return val if isinstance(val, dict) else None
    except json.JSONDecodeError:
        return None


def _batch_index_from_job(job: SystemBgJobQueue, payload: dict) -> int:
    if payload.get("batch_index") is not None:
        return int(payload["batch_index"])
    result = _parse_json_field(job.result)
    if result and result.get("batch_index") is not None:
        return int(result["batch_index"])
    return 0


def _import_jobs_for_dataset(
    db: Session,
    *,
    dataset_id: int,
    dw_id: int,
    request_id: Optional[str] = None,
) -> list[SystemBgJobQueue]:
    q = db.query(SystemBgJobQueue).filter(
        SystemBgJobQueue.dataset_id == dataset_id,
        SystemBgJobQueue.dw_id == str(dw_id),
        SystemBgJobQueue.job_type == JOB_TYPE_IMPORT_DATASET_DATA_POINTS_BATCH,
    )
    if request_id is not None:
        q = q.filter(SystemBgJobQueue.request_id == request_id)
    return q.all()


def list_import_logs_for_dataset(
    db: Session,
    *,
    dataset_id: int,
    dw_id: int,
) -> dict[str, Any]:
    assert_import_allowed(db, dataset_id=dataset_id, dw_id=dw_id, check_readiness=False, check_editable=False)

    jobs = _import_jobs_for_dataset(db, dataset_id=dataset_id, dw_id=dw_id)
    by_request: dict[str, list[SystemBgJobQueue]] = defaultdict(list)
    for job in jobs:
        if job.request_id:
            by_request[job.request_id].append(job)

    items: list[dict[str, Any]] = []
    for rid, job_list in by_request.items():
        summary = _summarize_request_jobs(job_list)
        summary["request_id"] = rid
        items.append(summary)

    items.sort(
        key=lambda x: (
            x.get("finished_at") is None,
            -(x.get("sort_ts") or 0),
        ),
    )
    for it in items:
        it.pop("sort_ts", None)

    return {"dataset_id": dataset_id, "imports": items}


def get_import_request_detail(
    db: Session,
    *,
    dataset_id: int,
    dw_id: int,
    request_id: str,
) -> dict[str, Any]:
    assert_import_allowed(db, dataset_id=dataset_id, dw_id=dw_id, check_readiness=False, check_editable=False)

    jobs = _import_jobs_for_dataset(
        db, dataset_id=dataset_id, dw_id=dw_id, request_id=request_id
    )
    if not jobs:
        raise ValueError("IMPORT_REQUEST_NOT_FOUND")

    jobs_sorted = sorted(
        jobs,
        key=lambda j: _batch_index_from_job(j, _parse_job_payload(j)),
    )

    batches: list[dict[str, Any]] = []
    for job in jobs_sorted:
        payload = _parse_job_payload(job)
        batches.append({
            "job_id": job.id,
            "batch_index": _batch_index_from_job(job, payload),
            "job_status": job.job_status,
            "result": _parse_json_field(job.result),
            "last_error": job.last_error,
            "completed_at": job.completed_at,
            "created_at": job.created_at,
        })

    summary = _summarize_request_jobs(jobs)
    summary.pop("sort_ts", None)
    return {
        "dataset_id": dataset_id,
        "request_id": request_id,
        **summary,
        "batches": batches,
    }


def _parse_job_payload(job: SystemBgJobQueue) -> dict:
    try:
        return json.loads(job.job_req_payload)
    except json.JSONDecodeError:
        return {}


def _summarize_request_jobs(jobs: list[SystemBgJobQueue]) -> dict[str, Any]:
    total_batches = len(jobs)
    completed_batches = 0
    failed_batches = 0
    pending_batches = 0
    total_updated = 0
    total_skipped_locked = 0
    started_at = None
    finished_at = None
    sort_ts = 0.0

    for job in jobs:
        if job.created_at and (started_at is None or job.created_at < started_at):
            started_at = job.created_at
        if job.completed_at:
            if finished_at is None or job.completed_at > finished_at:
                finished_at = job.completed_at
            ts = job.completed_at.timestamp()
            if ts > sort_ts:
                sort_ts = ts
        elif job.created_at:
            ts = job.created_at.timestamp()
            if ts > sort_ts:
                sort_ts = ts

        if job.job_status == "completed":
            completed_batches += 1
            result = _parse_json_field(job.result)
            if result:
                total_updated += int(result.get("updated") or 0)
                total_skipped_locked += int(result.get("skipped_locked") or 0)
        elif job.job_status == "failed":
            failed_batches += 1
        elif job.job_status in _ACTIVE_JOB_STATUSES:
            pending_batches += 1

    if pending_batches > 0:
        overall_status = "running"
    elif failed_batches > 0:
        overall_status = "failed"
    else:
        overall_status = "completed"

    return {
        "started_at": started_at,
        "finished_at": finished_at,
        "sort_ts": sort_ts,
        "total_batches": total_batches,
        "completed_batches": completed_batches,
        "failed_batches": failed_batches,
        "pending_batches": pending_batches,
        "total_updated": total_updated,
        "total_skipped_locked": total_skipped_locked,
        "overall_status": overall_status,
    }
