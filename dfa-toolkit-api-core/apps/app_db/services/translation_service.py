"""
Service for translations: lookup locale strings from system_language_translations
and system_language_translations_i18n.
"""
from io import BytesIO
from datetime import datetime
from typing import Any, Optional

from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter
from sqlalchemy import or_
from sqlalchemy.orm import Session

from apps.app_db.models.core import (
    SystemApplication,
    SystemLanguage,
    SystemLanguageTranslation,
    SystemLanguageTranslationI18n,
)
from apps.app_db.db.core import SessionLocal
from apps.shared.exceptions import Conflict
from apps.app_db.schemas.core import (
    SystemLanguageTranslationCreateSchema,
    SystemLanguageTranslationPatchSchema,
)
from apps.shared.utils.log_event import log_background_incident

_TRANSLATION_SERVICE_SOURCE = "apps.app_db.services.translation_service"


def _translation_list_search_matches(
    trans_code: str,
    default_text: Optional[str],
    locale_texts: dict[str, str],
    search: Optional[str],
    search_by: Optional[str],
) -> bool:
    """Case-insensitive substring filter for merged translation list/export rows."""
    needle = (search or "").strip().lower()
    if not needle:
        return True
    if not search_by:
        return True
    dt = (default_text or "").lower()
    tc = trans_code.lower()
    if search_by == "trans_code":
        return needle in tc
    if search_by == "default_text":
        return needle in dt
    if search_by == "translation_text":
        for v in locale_texts.values():
            if needle in (v or "").lower():
                return True
        return needle in dt
    return True


def _sort_translation_list_rows(
    rows: list[tuple[str, Any, dict[str, str]]],
    sort_field: str,
    sort_order: str,
    first_locale_for_text_sort: str,
) -> list[tuple[str, Any, dict[str, str]]]:
    """Sort merged rows by sort_field; first_locale_for_text_sort drives translation_text key."""
    reverse = (sort_order or "").strip().lower() == "desc"

    def sort_key(item: tuple[str, Any, dict[str, str]]) -> tuple:
        code, master, lt = item
        mid = int(master.id)
        dt = master.default_text or ""
        first_txt = lt.get(first_locale_for_text_sort)
        if first_txt is None:
            first_txt = dt
        if sort_field == "id":
            return (mid, code.lower())
        if sort_field == "default_text":
            return (dt.lower(), mid)
        if sort_field == "translation_text":
            return ((first_txt or "").lower(), mid)
        return (code.lower(), mid)

    return sorted(rows, key=sort_key, reverse=reverse)


class TranslationService:
    """
    Service for translation lookup:
    - Get translations for a locale (app_key IS NULL for global, or NULL + app_key).
    - Resolve app_key (keyword) to app_id.
    """

    def resolve_app_key_to_app_id(self, keyword: str) -> Optional[int]:
        """
        Resolve application keyword (app_key) to system_applications.id.

        Args:
            keyword: Application key (app_key).

        Returns:
            Application id, or None if no application has that app_key.
        """
        db: Session = SessionLocal()
        try:
            app = (
                db.query(SystemApplication)
                .filter(SystemApplication.app_key == keyword)
                .first()
            )
            return app.id if app else None
        finally:
            db.close()

    def get_translations_for_locale(
        self,
        locale: str,
        app_key_filter: Optional[str] = None,
    ) -> dict[str, str]:
        """
        Get translation code -> text for the given locale.

        - If app_key_filter is None: only rows with app_key IS NULL (global).
        - If app_key_filter is set: rows with app_key IS NULL or app_key = app_key_filter,
          merged with global first, then app-specific overrides for the same code.

        For each master row, use system_language_translations_i18n.translated_text
        when (master_id, locale) exists, else system_language_translations.default_text.
        Rows with code IS NULL are skipped.

        Args:
            locale: Normalized locale code (e.g. "en", "fr").
            app_key_filter: If set, include global (app_key NULL) and this app_key; app-specific overrides global.

        Returns:
            Dict mapping code -> translated or default text.
        """
        db: Session = SessionLocal()
        try:
            if app_key_filter is None:
                # Global only: app_key IS NULL
                app_filter = SystemLanguageTranslation.app_key.is_(None)
            else:
                # Global (NULL) + app-specific; NULL first so app-specific overrides in merge
                app_filter = or_(
                    SystemLanguageTranslation.app_key.is_(None),
                    SystemLanguageTranslation.app_key == app_key_filter,
                )

            # ORDER BY app_key: MySQL puts NULL first in ASC, so global then app-specific
            masters = (
                db.query(SystemLanguageTranslation)
                .filter(app_filter)
                .order_by(SystemLanguageTranslation.app_key)
                .all()
            )
            if not masters:
                return {}

            master_ids = [m.id for m in masters]

            # Bulk load i18n for these master_ids and locale
            i18n_rows = (
                db.query(
                    SystemLanguageTranslationI18n.master_id,
                    SystemLanguageTranslationI18n.translated_text,
                )
                .filter(
                    SystemLanguageTranslationI18n.master_id.in_(master_ids),
                    SystemLanguageTranslationI18n.locale == locale,
                )
                .all()
            )
            i18n_map = {r.master_id: r.translated_text for r in i18n_rows}

            # Build code -> text; later rows (app-specific) override earlier (global)
            result: dict[str, str] = {}
            for m in masters:
                if m.trans_code is None:
                    continue
                text = i18n_map.get(m.id)
                if text is None:
                    text = m.default_text
                result[m.trans_code] = text

            return result
        finally:
            db.close()

    def get_translation_export_table(
        self,
        app_key_filter: Optional[str] = None,
        search: Optional[str] = None,
        search_by: Optional[str] = None,
        sort_field: str = "trans_code",
        sort_order: str = "asc",
    ) -> dict[str, list[dict[str, str]] | list[str]]:
        """
        Export translations in spreadsheet-like table format.

        Returns:
            {
                "columns": ["Code", "<locale1>", "<locale2>", ...],
                "rows": [
                    {"Code": "ABOUT", "en": "About", "fr": "A propos", ...},
                    ...
                ]
            }
        """
        db: Session = SessionLocal()
        try:
            if app_key_filter is None:
                app_filter = SystemLanguageTranslation.app_key.is_(None)
            else:
                app_filter = or_(
                    SystemLanguageTranslation.app_key.is_(None),
                    SystemLanguageTranslation.app_key == app_key_filter,
                )

            locales = [
                row.locale
                for row in (
                    db.query(SystemLanguage)
                    .filter(SystemLanguage.is_deleted.is_(False))
                    .order_by(SystemLanguage.is_default.desc(), SystemLanguage.locale.asc())
                    .all()
                )
                if row.locale
            ]

            masters = (
                db.query(SystemLanguageTranslation)
                .filter(app_filter)
                .order_by(SystemLanguageTranslation.app_key, SystemLanguageTranslation.id)
                .all()
            )
            if not masters:
                return {"columns": ["Code", *locales], "rows": []}

            effective_masters_by_code: dict[str, SystemLanguageTranslation] = {}
            for master in masters:
                if not master.trans_code:
                    continue
                # For app export, app-specific rows come after global rows and override.
                effective_masters_by_code[master.trans_code] = master

            if not effective_masters_by_code:
                return {"columns": ["Code", *locales], "rows": []}

            master_ids = [m.id for m in effective_masters_by_code.values()]
            i18n_query = db.query(
                SystemLanguageTranslationI18n.master_id,
                SystemLanguageTranslationI18n.locale,
                SystemLanguageTranslationI18n.translated_text,
            ).filter(SystemLanguageTranslationI18n.master_id.in_(master_ids))
            if locales:
                i18n_query = i18n_query.filter(SystemLanguageTranslationI18n.locale.in_(locales))
            i18n_rows = i18n_query.all()
            i18n_map = {(r.master_id, r.locale): r.translated_text for r in i18n_rows}

            row_tuples: list[tuple[str, SystemLanguageTranslation, dict[str, str]]] = []
            for code, master in effective_masters_by_code.items():
                locale_texts: dict[str, str] = {}
                for loc in locales:
                    locale_texts[loc] = i18n_map.get((master.id, loc), master.default_text or "")
                row_tuples.append((code, master, locale_texts))

            filtered = [
                t
                for t in row_tuples
                if _translation_list_search_matches(
                    t[0], t[1].default_text, t[2], search, search_by
                )
            ]
            first_loc = locales[0] if locales else ""
            sorted_rows = _sort_translation_list_rows(
                filtered, sort_field, sort_order, first_loc
            )

            rows: list[dict[str, str]] = []
            for code, master, locale_texts in sorted_rows:
                row_data: dict[str, str] = {"Code": code}
                for loc in locales:
                    row_data[loc] = locale_texts.get(loc, master.default_text or "")
                rows.append(row_data)

            return {"columns": ["Code", *locales], "rows": rows}
        finally:
            db.close()

    def get_translation_export_xlsx(
        self,
        app_key_filter: Optional[str] = None,
        search: Optional[str] = None,
        search_by: Optional[str] = None,
        sort_field: str = "trans_code",
        sort_order: str = "asc",
    ) -> tuple[BytesIO, str]:
        """
        Build XLSX for translation export and return stream + filename.
        """
        table = self.get_translation_export_table(
            app_key_filter=app_key_filter,
            search=search,
            search_by=search_by,
            sort_field=sort_field,
            sort_order=sort_order,
        )
        columns = table.get("columns", [])
        rows = table.get("rows", [])

        wb = Workbook()
        ws = wb.active
        ws.title = "Translations"
        ws.append(columns)

        max_lens = [len(str(col)) for col in columns]
        for row in rows:
            row_values = [row.get(col, "") for col in columns]
            ws.append(row_values)
            for idx, val in enumerate(row_values):
                sval = "" if val is None else str(val)
                if idx < len(max_lens) and len(sval) > max_lens[idx]:
                    max_lens[idx] = len(sval)

        for i, max_len in enumerate(max_lens, start=1):
            ws.column_dimensions[get_column_letter(i)].width = min(max(max_len + 4, 12), 80)

        output = BytesIO()
        wb.save(output)
        output.seek(0)

        ts = datetime.now().strftime("%Y-%m-%d-%H-%M-%S")
        scope = (app_key_filter or "global").strip() if app_key_filter else "global"
        filename = f"TRANSLATIONS_{scope}_{ts}.xlsx"
        return output, filename

    @staticmethod
    def _normalize_excel_cell(value: Any) -> str:
        """Convert an Excel cell value to trimmed string."""
        if value is None:
            return ""
        return str(value).strip()

    @staticmethod
    def _normalize_app_key(app_key: Optional[str]) -> Optional[str]:
        if app_key is None:
            return None
        normalized = app_key.strip()
        if not normalized:
            return None
        return normalized

    def _get_unique_master_by_code(
        self,
        db: Session,
        code: str,
    ) -> Optional[SystemLanguageTranslation]:
        rows = db.query(SystemLanguageTranslation).filter(
            SystemLanguageTranslation.trans_code == code
        ).all()
        if not rows:
            return None
        if len(rows) > 1:
            raise ValueError(
                f"Multiple translations found for code '{code}'. Please clean duplicate data first."
            )
        return rows[0]

    def _get_master_by_code_and_scope(
        self,
        db: Session,
        code: str,
        app_key_filter: Optional[str],
    ) -> Optional[SystemLanguageTranslation]:
        query = db.query(SystemLanguageTranslation).filter(
            SystemLanguageTranslation.trans_code == code
        )
        if app_key_filter is None:
            query = query.filter(SystemLanguageTranslation.app_key.is_(None))
        else:
            query = query.filter(SystemLanguageTranslation.app_key == app_key_filter)
        rows = query.all()
        if not rows:
            return None
        if len(rows) > 1:
            scope = app_key_filter or "global"
            raise ValueError(
                f"Multiple translations found for code '{code}' in scope '{scope}'. "
                "Please clean duplicate data first."
            )
        return rows[0]

    def _get_master_by_code_app_then_global(
        self,
        db: Session,
        code: str,
        resolved_app_key: Optional[str],
    ) -> Optional[SystemLanguageTranslation]:
        """
        Resolve a master row for PATCH/DELETE: prefer app-scoped (resolved_app_key),
        then fall back to global (app_key IS NULL). Caller passes the app key from
        X-Api-Key resolution (may be None after normalize).
        """
        normalized = self._normalize_app_key(resolved_app_key)
        if normalized is not None:
            row = self._get_master_by_code_and_scope(db, code, normalized)
            if row is not None:
                return row
        return self._get_master_by_code_and_scope(db, code, None)

    @staticmethod
    def _get_active_locales(db: Session) -> set[str]:
        locales = (
            db.query(SystemLanguage.locale)
            .filter(SystemLanguage.is_deleted.is_(False))
            .all()
        )
        return {locale.strip().lower() for (locale,) in locales if locale and locale.strip()}

    @staticmethod
    def _get_active_languages(db: Session) -> list[SystemLanguage]:
        return (
            db.query(SystemLanguage)
            .filter(SystemLanguage.is_deleted.is_(False))
            .order_by(SystemLanguage.is_default.desc(), SystemLanguage.locale.asc())
            .all()
        )

    @staticmethod
    def _get_active_locales_ordered(db: Session) -> list[str]:
        locales: list[str] = []
        for language in TranslationService._get_active_languages(db):
            locale_key = (language.locale or "").strip().lower()
            if locale_key and locale_key not in locales:
                locales.append(locale_key)
        return locales

    def _validate_translation_locales(
        self,
        db: Session,
        translations: Optional[dict[str, str]],
    ) -> None:
        if not translations:
            return
        active_locales = self._get_active_locales(db)
        if not active_locales:
            raise ValueError(
                "No active system language found. Please create a system language first."
            )

        invalid_locales = sorted(
            {
                locale.strip().lower()
                for locale in translations.keys()
                if locale and locale.strip() and locale.strip().lower() not in active_locales
            }
        )
        if invalid_locales:
            raise ValueError(
                "Invalid locale(s): "
                f"{', '.join(invalid_locales)}. Allowed locales: {', '.join(sorted(active_locales))}"
            )

    def _get_effective_default_locale(self, db: Session) -> Optional[str]:
        """
        Resolve default locale from active languages.
        Falls back to the first active locale when no explicit default is set.
        """
        languages = self._get_active_languages(db)
        for language in languages:
            locale_key = (language.locale or "").strip().lower()
            if not locale_key:
                continue
            if language.is_default:
                return locale_key
        for language in languages:
            locale_key = (language.locale or "").strip().lower()
            if locale_key:
                return locale_key
        return None

    def _derive_default_text_from_translations(
        self,
        db: Session,
        translations: Optional[dict[str, str]],
    ) -> str:
        """
        Derive default_text from translations using:
        1) default locale text from system_language if present and non-empty
        2) first non-empty translation value
        """
        if not translations:
            raise ValueError(
                "translations must include at least one locale text to derive default_text"
            )

        normalized_pairs: list[tuple[str, str]] = []
        for locale, text in translations.items():
            locale_key = (locale or "").strip().lower()
            if not locale_key:
                continue
            normalized_pairs.append((locale_key, text))

        if not normalized_pairs:
            raise ValueError(
                "translations must include at least one valid locale key to derive default_text"
            )

        default_locale = self._get_effective_default_locale(db)
        if default_locale is not None:
            for locale_key, text in normalized_pairs:
                if locale_key == default_locale and str(text).strip() != "":
                    return text

        for _, text in normalized_pairs:
            if str(text).strip() != "":
                return text

        raise ValueError("translations must include at least one non-empty text value")

    def import_translations_from_xlsx(
        self,
        file_bytes: bytes,
        app_key_filter: Optional[str],
        user_id: Optional[int],
    ) -> dict[str, int | str]:
        """
        Import translations from XLSX where:
        - Column 1 is "Code"
        - Remaining columns are locale codes from system_language

        Behavior:
        - Skip rows with empty Code
        - Reject duplicate Code rows in file
        - Reject duplicate locale columns in file
        - Reject duplicate locale codes in system_language
        - Upsert rows in the provided app scope (global if app_key is blank/None)
        - Update existing i18n/default_text values when changed
        """
        db: Session = SessionLocal()
        try:
            if not file_bytes:
                raise ValueError("Uploaded Excel file is empty")

            workbook = load_workbook(filename=BytesIO(file_bytes), data_only=True)
            worksheet = workbook.active
            rows = worksheet.iter_rows(values_only=True)
            header = next(rows, None)
            if not header:
                raise ValueError("Excel file does not contain a header row")

            normalized_header = [self._normalize_excel_cell(v) for v in header]
            if not normalized_header:
                raise ValueError("Excel file header is empty")
            if normalized_header[0].lower() != "code":
                raise ValueError("First column must be 'Code'")

            locale_columns_raw = normalized_header[1:]
            locale_columns = [locale.strip().lower() for locale in locale_columns_raw if locale.strip()]
            if not locale_columns:
                raise ValueError("Excel must contain at least one locale column after 'Code'")

            duplicate_locale_columns = sorted(
                {
                    locale
                    for locale in locale_columns
                    if locale_columns.count(locale) > 1
                }
            )
            if duplicate_locale_columns:
                raise ValueError(
                    f"Duplicate locale column(s) in Excel: {', '.join(duplicate_locale_columns)}"
                )

            language_rows = self._get_active_languages(db)
            if not language_rows:
                raise ValueError("No active system language found. Please create a system language first.")

            locale_counts: dict[str, int] = {}
            default_locale: Optional[str] = None
            active_locales_ordered: list[str] = []
            for language in language_rows:
                locale_key = (language.locale or "").strip().lower()
                if not locale_key:
                    continue
                locale_counts[locale_key] = locale_counts.get(locale_key, 0) + 1
                if locale_key not in active_locales_ordered:
                    active_locales_ordered.append(locale_key)
                if language.is_default and default_locale is None:
                    default_locale = locale_key

            duplicate_system_locales = sorted(
                [locale for locale, count in locale_counts.items() if count > 1]
            )
            if duplicate_system_locales:
                raise ValueError(
                    "Duplicate locale code(s) found in system_language: "
                    f"{', '.join(duplicate_system_locales)}"
                )

            active_locale_set = set(active_locales_ordered)
            invalid_excel_locales = sorted(
                [locale for locale in locale_columns if locale not in active_locale_set]
            )
            if invalid_excel_locales:
                raise ValueError(
                    "Invalid locale column(s): "
                    f"{', '.join(invalid_excel_locales)}. "
                    f"Allowed locales: {', '.join(active_locales_ordered)}"
                )

            if default_locale is None:
                default_locale = active_locales_ordered[0]

            parsed_rows: list[dict[str, Any]] = []
            seen_codes: set[str] = set()
            duplicate_codes: set[str] = set()
            skipped_empty_code = 0
            for row in rows:
                values = list(row or [])
                code = self._normalize_excel_cell(values[0] if values else "")
                if not code:
                    skipped_empty_code += 1
                    continue

                code_key = code.lower()
                if code_key in seen_codes:
                    duplicate_codes.add(code)
                    continue
                seen_codes.add(code_key)

                translations: dict[str, str] = {}
                for idx, locale in enumerate(locale_columns, start=1):
                    cell_value = self._normalize_excel_cell(values[idx] if idx < len(values) else "")
                    translations[locale] = cell_value

                default_text_candidate = translations.get(default_locale)
                if not default_text_candidate and translations:
                    default_text_candidate = next(iter(translations.values()))
                if not default_text_candidate:
                    default_text_candidate = code

                parsed_rows.append(
                    {
                        "code": code,
                        "translations": translations,
                        "default_text": default_text_candidate,
                    }
                )

            if duplicate_codes:
                raise ValueError(
                    f"Duplicate translation code(s) in Excel: {', '.join(sorted(duplicate_codes))}"
                )

            if not parsed_rows:
                return {
                    "app_key": self._normalize_app_key(app_key_filter) or "global",
                    "total_rows": 0,
                    "inserted": 0,
                    "updated": 0,
                    "skipped_empty_code": skipped_empty_code,
                }

            normalized_app_key = self._normalize_app_key(app_key_filter)
            code_list = [item["code"] for item in parsed_rows]
            existing_query = db.query(SystemLanguageTranslation).filter(
                SystemLanguageTranslation.trans_code.in_(code_list)
            )
            if normalized_app_key is None:
                existing_query = existing_query.filter(SystemLanguageTranslation.app_key.is_(None))
            else:
                existing_query = existing_query.filter(
                    SystemLanguageTranslation.app_key == normalized_app_key
                )
            existing_rows = existing_query.all()

            existing_by_code: dict[str, SystemLanguageTranslation] = {}
            duplicate_db_codes: set[str] = set()
            for existing_row in existing_rows:
                existing_key = (existing_row.trans_code or "").strip().lower()
                if not existing_key:
                    continue
                if existing_key in existing_by_code:
                    duplicate_db_codes.add(existing_row.trans_code)
                    continue
                existing_by_code[existing_key] = existing_row

            if duplicate_db_codes:
                scope = normalized_app_key or "global"
                raise ValueError(
                    f"Duplicate translation code(s) already exist in '{scope}' scope: "
                    f"{', '.join(sorted(duplicate_db_codes))}"
                )

            existing_ids = [row.id for row in existing_by_code.values()]
            i18n_by_key: dict[tuple[int, str], SystemLanguageTranslationI18n] = {}
            if existing_ids:
                i18n_rows = (
                    db.query(SystemLanguageTranslationI18n)
                    .filter(
                        SystemLanguageTranslationI18n.master_id.in_(existing_ids),
                        SystemLanguageTranslationI18n.locale.in_(locale_columns),
                    )
                    .all()
                )
                i18n_by_key = {(row.master_id, row.locale): row for row in i18n_rows}

            inserted = 0
            updated = 0
            now = datetime.utcnow()
            for item in parsed_rows:
                code = item["code"]
                translations = item["translations"]
                default_text_value = item["default_text"]
                code_key = code.lower()
                master = existing_by_code.get(code_key)
                row_changed = False
                is_new_row = False

                if master is None:
                    master = SystemLanguageTranslation(
                        trans_code=code,
                        app_key=normalized_app_key,
                        default_text=default_text_value,
                        created_user_id=user_id or 0,
                        created=now,
                        modified_user_id=user_id,
                        modified=now,
                    )
                    db.add(master)
                    db.flush()
                    existing_by_code[code_key] = master
                    inserted += 1
                    row_changed = True
                    is_new_row = True
                else:
                    if default_text_value and master.default_text != default_text_value:
                        master.default_text = default_text_value
                        row_changed = True

                for locale, text in translations.items():
                    existing_i18n = i18n_by_key.get((master.id, locale))
                    if existing_i18n:
                        if existing_i18n.translated_text != text:
                            existing_i18n.translated_text = text
                            row_changed = True
                    else:
                        new_i18n = SystemLanguageTranslationI18n(
                            master_id=master.id,
                            locale=locale,
                            translated_text=text,
                        )
                        db.add(new_i18n)
                        i18n_by_key[(master.id, locale)] = new_i18n
                        row_changed = True

                if row_changed and not is_new_row:
                    master.modified_user_id = user_id
                    master.modified = now
                if row_changed and not is_new_row:
                    updated += 1

            db.commit()
            return {
                "app_key": normalized_app_key or "global",
                "total_rows": len(parsed_rows),
                "inserted": inserted,
                "updated": updated,
                "skipped_empty_code": skipped_empty_code,
            }
        except Exception:
            db.rollback()
            raise
        finally:
            db.close()

    def get_translations_for_locale_paginated(
        self,
        locale: str,
        app_key_filter: Optional[str] = None,
        page: int = 1,
        page_size: Optional[int] = None,
        requested_locale: Optional[str] = None,
        search: Optional[str] = None,
        search_by: Optional[str] = None,
        sort_field: str = "trans_code",
        sort_order: str = "asc",
    ) -> tuple[list[dict[str, dict[str, str]]], int, Optional[int], Optional[int]]:
        """
        Return translations as paginated list of code keyed locale maps.

        Response item shape:
        - requested locale mode: {"ABOUT": {"en": "About", "fr": "A propos"}}
        - all locale mode: {"ABOUT": {"en": "About", "fr": "A propos", "ar": "..."}}

        Optional ``search`` / ``search_by`` filter rows before pagination;
        ``sort_field`` / ``sort_order`` order rows before pagination.
        """
        db: Session = SessionLocal()
        try:
            if app_key_filter is None:
                app_filter = SystemLanguageTranslation.app_key.is_(None)
            else:
                app_filter = or_(
                    SystemLanguageTranslation.app_key.is_(None),
                    SystemLanguageTranslation.app_key == app_key_filter,
                )

            active_locales = self._get_active_locales_ordered(db)
            if not active_locales:
                raise ValueError(
                    "No active system language found. Please create a system language first."
                )
            # system default locale is first in ordered active locales
            system_default_locale = active_locales[0]
            from apps.shared.utils.validators import normalize_language_code
            accepted_locale = normalize_language_code((locale or "").strip())
            default_locale = accepted_locale if accepted_locale in active_locales else system_default_locale

            requested_locale_value = (requested_locale or "").strip().lower()
            if not requested_locale_value:
                requested_locale_value = default_locale

            if requested_locale_value == "all":
                selected_locales = active_locales
            else:
                normalized_requested = normalize_language_code(requested_locale_value)
                if normalized_requested not in active_locales:
                    raise ValueError(
                        "Invalid locale(s): "
                        f"{normalized_requested}. Allowed locales: {', '.join(active_locales)}, all"
                    )
                selected_locales = [default_locale]
                if normalized_requested != default_locale:
                    selected_locales.append(normalized_requested)

            masters = (
                db.query(SystemLanguageTranslation)
                .filter(app_filter)
                .order_by(SystemLanguageTranslation.app_key, SystemLanguageTranslation.id)
                .all()
            )
            if not masters:
                return [], 0, page_size, None

            effective_masters_by_code: dict[str, SystemLanguageTranslation] = {}
            for master in masters:
                if not master.trans_code:
                    continue
                effective_masters_by_code[master.trans_code] = master

            if not effective_masters_by_code:
                return [], 0, page_size, None

            master_ids = [m.id for m in effective_masters_by_code.values()]
            i18n_rows = (
                db.query(
                    SystemLanguageTranslationI18n.master_id,
                    SystemLanguageTranslationI18n.locale,
                    SystemLanguageTranslationI18n.translated_text,
                )
                .filter(
                    SystemLanguageTranslationI18n.master_id.in_(master_ids),
                    SystemLanguageTranslationI18n.locale.in_(selected_locales),
                )
                .all()
            )
            i18n_map = {(row.master_id, row.locale): row.translated_text for row in i18n_rows}

            row_tuples: list[tuple[str, SystemLanguageTranslation, dict[str, str]]] = []
            for code, master in effective_masters_by_code.items():
                locale_texts: dict[str, str] = {}
                for locale_key in selected_locales:
                    locale_texts[locale_key] = i18n_map.get(
                        (master.id, locale_key), master.default_text
                    )
                row_tuples.append((code, master, locale_texts))

            has_list_filters = bool(search and search.strip())
            total_count_unfiltered: Optional[int] = (
                len(row_tuples) if has_list_filters else None
            )
            filtered = [
                t
                for t in row_tuples
                if _translation_list_search_matches(
                    t[0], t[1].default_text, t[2], search, search_by
                )
            ]
            first_loc_for_sort = selected_locales[0] if selected_locales else ""
            sorted_rows = _sort_translation_list_rows(
                filtered, sort_field, sort_order, first_loc_for_sort
            )

            records: list[dict[str, dict[str, str]]] = [
                {code: locale_texts} for code, _master, locale_texts in sorted_rows
            ]

            total_count = len(records)
            if page_size is None:
                return records, total_count, None, total_count_unfiltered

            effective_page_size = max(1, min(page_size, 1000))
            skip = (page - 1) * effective_page_size
            return (
                records[skip: skip + effective_page_size],
                total_count,
                effective_page_size,
                total_count_unfiltered,
            )
        finally:
            db.close()

    def create_translation(
        self,
        payload: SystemLanguageTranslationCreateSchema,
        user_id: Optional[int],
        app_key_filter: Optional[str] = None,
    ) -> SystemLanguageTranslation:
        db: Session = SessionLocal()
        try:
            normalized_app_key = self._normalize_app_key(app_key_filter)
            code = payload.code.strip()

            existing = self._get_unique_master_by_code(db, code)
            if existing:
                raise Conflict(f"Translation code '{code}' already exists")

            self._validate_translation_locales(db=db, translations=payload.translations)
            derived_default_text = self._derive_default_text_from_translations(
                db=db,
                translations=payload.translations,
            )

            now = datetime.utcnow()
            row = SystemLanguageTranslation(
                trans_code=code,
                app_key=normalized_app_key,
                default_text=derived_default_text,
                created_user_id=user_id or 0,
                created=now,
                modified_user_id=user_id,
                modified=now,
            )
            db.add(row)
            db.flush()

            for locale, text in payload.translations.items():
                locale_key = locale.strip().lower()
                if not locale_key:
                    continue
                db.add(
                    SystemLanguageTranslationI18n(
                        master_id=row.id,
                        locale=locale_key,
                        translated_text=text,
                    )
                )

            db.commit()
            db.refresh(row)
            return row
        except Exception:
            db.rollback()
            raise
        finally:
            db.close()

    def patch_translation(
        self,
        translation_code: str,
        payload: SystemLanguageTranslationPatchSchema,
        user_id: Optional[int],
        resolved_app_key: Optional[str] = None,
    ) -> Optional[SystemLanguageTranslation]:
        db: Session = SessionLocal()
        try:
            row = self._get_master_by_code_app_then_global(
                db=db,
                code=translation_code,
                resolved_app_key=resolved_app_key,
            )
            if not row:
                return None

            now = datetime.utcnow()
            # Keep translation_code rename logic commented for future enablement.
            # if payload.translation_code is not None:
            #     updated_code = payload.translation_code.strip()
            #     if not updated_code:
            #         raise ValueError("translation_code cannot be blank")
            #     if updated_code != row.trans_code:
            #         existing = self._get_unique_master_by_code(db, updated_code)
            #         if existing:
            #             raise ValueError(f"Translation code '{updated_code}' already exists")
            #         row.trans_code = updated_code

            if payload.translations:
                self._validate_translation_locales(db=db, translations=payload.translations)
                default_locale = self._get_effective_default_locale(db)
                default_locale_provided = False
                default_locale_text: Optional[str] = None
                for locale, text in payload.translations.items():
                    locale_key = locale.strip().lower()
                    if not locale_key:
                        continue
                    if default_locale is not None and locale_key == default_locale:
                        default_locale_provided = True
                        default_locale_text = text
                    i18n_row = db.query(SystemLanguageTranslationI18n).filter(
                        SystemLanguageTranslationI18n.master_id == row.id,
                        SystemLanguageTranslationI18n.locale == locale_key,
                    ).first()
                    if i18n_row:
                        i18n_row.translated_text = text
                    else:
                        db.add(
                            SystemLanguageTranslationI18n(
                                master_id=row.id,
                                locale=locale_key,
                                translated_text=text,
                            )
                        )
                if default_locale_provided:
                    row.default_text = default_locale_text or ""

            row.modified_user_id = user_id
            row.modified = now
            db.commit()
            db.refresh(row)
            return row
        except Exception:
            db.rollback()
            raise
        finally:
            db.close()

    def get_translation_by_code(
        self,
        translation_code: str,
    ) -> Optional[dict]:
        db: Session = SessionLocal()
        try:
            row = self._get_unique_master_by_code(db, translation_code)
            if not row:
                return None

            i18n_rows = (
                db.query(SystemLanguageTranslationI18n)
                .filter(SystemLanguageTranslationI18n.master_id == row.id)
                .all()
            )
            translations = {r.locale: r.translated_text for r in i18n_rows if r.locale}

            return {
                "id": row.id,
                "code": row.trans_code,
                "app_key": row.app_key,
                "default_text": row.default_text,
                "created_user_id": row.created_user_id,
                "created": row.created,
                "modified_user_id": row.modified_user_id,
                "modified": row.modified,
                "translations": translations,
            }
        finally:
            db.close()

    def delete_translation(
        self,
        translation_code: str,
        resolved_app_key: Optional[str] = None,
    ) -> bool:
        db: Session = SessionLocal()
        try:
            row = self._get_master_by_code_app_then_global(
                db, translation_code, resolved_app_key
            )
            if not row:
                return False

            db.query(SystemLanguageTranslationI18n).filter(
                SystemLanguageTranslationI18n.master_id == row.id
            ).delete(synchronize_session=False)
            db.delete(row)
            db.commit()
            return True
        except Exception:
            db.rollback()
            raise
        finally:
            db.close()

    def multiple_delete_translations(
        self,
        translation_codes: list[str],
        resolved_app_key: Optional[str] = None,
    ) -> dict[str, Any]:
        """
        Delete multiple translation masters (with i18n cascade) using the same
        app-then-global resolution as delete_translation. Partial success: each
        code is processed in its own transaction via delete_translation.
        """
        seen: set[str] = set()
        ordered_unique: list[str] = []
        for raw in translation_codes or []:
            code = (raw or "").strip()
            if not code or code in seen:
                continue
            seen.add(code)
            ordered_unique.append(code)

        deleted: list[str] = []
        errors: list[dict[str, Any]] = []

        for idx, code in enumerate(ordered_unique):
            try:
                ok = self.delete_translation(
                    translation_code=code,
                    resolved_app_key=resolved_app_key,
                )
                if ok:
                    deleted.append(code)
                else:
                    errors.append(
                        {
                            "index": idx,
                            "translation_code": code,
                            "error": f"Translation '{code}' not found",
                        }
                    )
            except ValueError as e:
                errors.append(
                    {
                        "index": idx,
                        "translation_code": code,
                        "error": str(e),
                    }
                )
            except Exception as e:
                log_background_incident(
                    message=f"multiple_delete_translations failed for code={code}",
                    error_code="TRANSLATION_MULTIPLE_DELETE_FAILED",
                    exc=e,
                    source=_TRANSLATION_SERVICE_SOURCE,
                )
                errors.append(
                    {
                        "index": idx,
                        "translation_code": code,
                        "error": str(e),
                    }
                )

        return {"deleted": deleted, "errors": errors}
