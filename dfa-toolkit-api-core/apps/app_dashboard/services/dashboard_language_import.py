"""
Dashboard language import — apply single-sheet XLSX from dashboard language export
onto edit-dashboard section names, header/footer HTML, and v6 visualization
property leaves as locale maps ``{lang_code: text}``.
"""
from __future__ import annotations

import json
import logging
import re
from datetime import datetime
from io import BytesIO
from typing import Any, Optional

from openpyxl import load_workbook
from sqlalchemy.orm import Session

from apps.app_dashboard.constants import DASHBOARD_MODE_EDIT
from apps.app_dashboard.models.dashboard import (
    Dashboard,
    DashboardSection,
    DashVisualization,
)
from apps.app_dashboard.services.dashboard_language_export import _load_language_codes
from apps.app_dashboard.services.section_service import SECTION_NAME_MAX_LENGTH
from apps.app_dashboard.utils.dashboard_serialization import (
    parse_json_field,
    serialize_json_column,
)
from apps.shared.exceptions import BadRequest, NotFound

logger = logging.getLogger(__name__)

_LANG_COLS_START = 5  # Excel column E (1-based) — matches export
_X000D_RE = re.compile(r"_x000D_", re.IGNORECASE)
_SERIES_PROP_RE = re.compile(r"^series\.(.+)$")
_CATEGORIES_PROP_RE = re.compile(r"^categories\.(.+)$")

# Property Name (col C) → nested path under VisualizationSettings / FilterSettings.
# Leaf is replaced with a locale map.
_GENERAL_TEXT_PATHS: dict[str, tuple[str, ...]] = {
    "noDataText": ("general", "chartArea", "noDataText"),
    "title": ("general", "title", "title", "text"),
    "subtitle": ("general", "title", "subTitle", "text"),
    "customLabel": ("general", "title", "customLabels", "text"),
    "xAxisTitle": ("general", "axisHorizontal", "title"),
    "yAxisTitle": ("general", "axisVertical", "title"),
    "htmlText": ("customHtml", "html"),
    "htmlNoDataText": ("customHtml", "noDataText"),
    "customText": ("customText", "text"),
    "filterLabel": ("label",),
}

_DASHBOARD_HTML_LABELS = {
    "Header HTML": "header_html",
    "Mobile Header HTML": "mobile_header_html",
    "Footer HTML": "footer_html",
}


def _cell_str(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _scrub_excel_artifacts(payload: str) -> str:
    return _X000D_RE.sub("", payload)


def _read_locale_map(
    ws,
    row: int,
    lang_col_by_code: dict[str, int],
) -> dict[str, str]:
    """Build ``{lang_code: cell}`` from language columns on ``row``."""
    result: dict[str, str] = {}
    for code, col in lang_col_by_code.items():
        result[code] = _cell_str(ws.cell(row=row, column=col).value)
    return result


def _set_nested(props: dict[str, Any], path: tuple[str, ...], value: Any) -> None:
    """Ensure intermediate dicts exist and set the leaf."""
    if not path:
        return
    cur: dict[str, Any] = props
    for key in path[:-1]:
        nxt = cur.get(key)
        if not isinstance(nxt, dict):
            nxt = {}
            cur[key] = nxt
        cur = nxt
    cur[path[-1]] = value


def _set_list_item_name(
    props: dict[str, Any],
    block_key: str,
    item_id: str,
    locale_map: dict[str, str],
) -> bool:
    """
    Set ``properties.{block_key}.{block_key}[].name`` for the item whose
    ``id`` (or index fallback) matches ``item_id``.
    """
    block = props.get(block_key)
    if not isinstance(block, dict):
        block = {}
        props[block_key] = block
    items = block.get(block_key)
    if not isinstance(items, list):
        items = []
        block[block_key] = items

    for idx, item in enumerate(items):
        if not isinstance(item, dict):
            continue
        key = str(item.get("id", item.get("index", idx)))
        if key == item_id:
            item["name"] = locale_map
            return True

    # Item missing from current properties — append so import is not silent.
    items.append({"id": item_id, "name": locale_map})
    return True


def apply_v6_viz_i18n_field(
    props: dict[str, Any],
    prop_label: str,
    key: Optional[str],
    locale_map: dict[str, str],
) -> bool:
    """
    Apply one Excel visualization row onto ``props``.

    Returns True when the property was recognized and written.
    """
    label = (prop_label or "").strip()
    if not label:
        return False

    path = _GENERAL_TEXT_PATHS.get(label)
    if path is not None:
        _set_nested(props, path, locale_map)
        return True

    series_m = _SERIES_PROP_RE.match(label)
    if series_m:
        return _set_list_item_name(props, "series", series_m.group(1), locale_map)

    cat_m = _CATEGORIES_PROP_RE.match(label)
    if cat_m:
        return _set_list_item_name(props, "categories", cat_m.group(1), locale_map)

    # Unknown property label — ignore (forward-compat with older sheets).
    logger.warning("Skipping unknown dashboard language property %r (key=%r)", label, key)
    return False


def _parse_lang_header(ws, dw_lang_codes: list[str]) -> dict[str, int]:
    """
    Map language code → 1-based column from row 1 starting at column E.

    Only codes present in both the sheet header and the warehouse language
    list are accepted.
    """
    mapping: dict[str, int] = {}
    max_col = ws.max_column or _LANG_COLS_START
    for col in range(_LANG_COLS_START, max_col + 1):
        raw = _cell_str(ws.cell(row=1, column=col).value)
        if not raw:
            continue
        # Prefer exact match against DW codes (case-insensitive).
        matched = next((c for c in dw_lang_codes if c.lower() == raw.lower()), None)
        if matched is None:
            logger.warning("Ignoring Excel language column %r (not in warehouse)", raw)
            continue
        mapping[matched] = col
    if not mapping:
        raise BadRequest(
            "Excel header has no language columns matching warehouse languages",
            error_code="DASHBOARD_LANG_IMPORT_NO_LANG_COLS",
            details={"warehouse_languages": dw_lang_codes},
        )
    return mapping


def _serialize_section_name(locale_map: dict[str, str]) -> str:
    serialized = json.dumps(locale_map, separators=(",", ":"))
    if len(serialized) > SECTION_NAME_MAX_LENGTH:
        raise BadRequest(
            f"section_name exceeds maximum length of {SECTION_NAME_MAX_LENGTH} characters",
            error_code="SECTION_NAME_TOO_LONG",
        )
    return serialized


def import_dashboard_languages_from_xlsx(
    core_db: Session,
    dw_db: Session,
    *,
    dashboard_id: int,
    system_dw_id: int,
    file_bytes: bytes,
) -> dict[str, Any]:
    """
    Import dashboard translation workbook produced by
    ``build_dashboard_languages_export_xlsx``.

    Writes locale maps onto v6 nested property leaves and dashboard HTML /
    section name JSON columns. Commits on the core session.
    """
    if not file_bytes:
        raise BadRequest(
            "Uploaded Excel file is empty",
            error_code="DASHBOARD_LANG_IMPORT_EMPTY_FILE",
        )

    dashboard = core_db.query(Dashboard).filter(Dashboard.id == dashboard_id).first()
    if not dashboard:
        raise NotFound(
            f"Dashboard with id {dashboard_id} not found",
            error_code="DASHBOARD_NOT_FOUND",
        )
    if dashboard.mode != DASHBOARD_MODE_EDIT:
        raise BadRequest(
            "Only edit-mode dashboards can be imported",
            error_code="DASHBOARD_NOT_EDIT_MODE",
            details={"dashboard_id": dashboard_id, "mode": dashboard.mode},
        )
    if dashboard.system_dw_id != system_dw_id:
        raise BadRequest(
            "Dashboard is not linked to the requested data warehouse",
            error_code="DASHBOARD_DW_MISMATCH",
            details={
                "dashboard_id": dashboard_id,
                "dashboard_system_dw_id": dashboard.system_dw_id,
                "path_system_dw_id": system_dw_id,
            },
        )

    lang_codes, _default_code = _load_language_codes(dw_db)
    if not lang_codes:
        raise BadRequest(
            "No languages configured for this data warehouse",
            error_code="DW_LANGUAGES_EMPTY",
        )

    try:
        wb = load_workbook(filename=BytesIO(file_bytes), data_only=True)
    except Exception as exc:
        raise BadRequest(
            "Invalid or unreadable Excel file",
            error_code="DASHBOARD_LANG_IMPORT_INVALID_XLSX",
            details={"reason": str(exc)},
        ) from exc

    ws = wb.active
    lang_col_by_code = _parse_lang_header(ws, lang_codes)

    sections = (
        core_db.query(DashboardSection)
        .filter(DashboardSection.dashboard_id == dashboard_id)
        .order_by(DashboardSection.order.asc(), DashboardSection.id.asc())
        .all()
    )
    section_iter = iter(sections)

    visualizations = (
        core_db.query(DashVisualization)
        .filter(
            DashVisualization.dashboard_id == dashboard_id,
            DashVisualization.is_deleted == 0,
        )
        .all()
    )
    viz_by_id: dict[int, DashVisualization] = {v.id: v for v in visualizations}
    # Accumulate property updates per viz so we write once.
    props_by_viz: dict[int, dict[str, Any]] = {}

    sections_updated = 0
    html_updated = 0
    viz_fields_updated = 0
    viz_rows_skipped = 0
    max_row = ws.max_row or 1

    for row in range(2, max_row + 1):
        col_a = _cell_str(ws.cell(row=row, column=1).value)
        if not col_a:
            continue

        if col_a == "Section Names":
            try:
                section = next(section_iter)
            except StopIteration:
                logger.warning(
                    "Extra Section Names row at Excel row %s; no matching section",
                    row,
                )
                continue
            locale_map = _read_locale_map(ws, row, lang_col_by_code)
            section.section_name = _serialize_section_name(locale_map)
            section.modified = datetime.utcnow()
            sections_updated += 1
            continue

        html_attr = _DASHBOARD_HTML_LABELS.get(col_a)
        if html_attr is not None:
            locale_map = _read_locale_map(ws, row, lang_col_by_code)
            raw = serialize_json_column(locale_map) or "{}"
            setattr(dashboard, html_attr, _scrub_excel_artifacts(raw))
            html_updated += 1
            continue

        if col_a == "Visualization IDs":
            continue

        # Visualization property row: col A = viz id, C = property, D = key
        try:
            viz_id = int(float(col_a))
        except (TypeError, ValueError):
            continue

        viz = viz_by_id.get(viz_id)
        if viz is None:
            viz_rows_skipped += 1
            logger.warning(
                "Skipping viz language row for unknown visualization id=%s",
                viz_id,
            )
            continue

        prop_label = _cell_str(ws.cell(row=row, column=3).value)
        key = _cell_str(ws.cell(row=row, column=4).value) or None
        locale_map = _read_locale_map(ws, row, lang_col_by_code)

        if viz_id not in props_by_viz:
            parsed = parse_json_field(viz.properties) or {}
            props_by_viz[viz_id] = parsed if isinstance(parsed, dict) else {}

        applied = apply_v6_viz_i18n_field(
            props_by_viz[viz_id],
            prop_label,
            key,
            locale_map,
        )
        if applied:
            viz_fields_updated += 1
        else:
            viz_rows_skipped += 1

    now = datetime.utcnow()
    for viz_id, props in props_by_viz.items():
        viz = viz_by_id[viz_id]
        payload = serialize_json_column(props) or "{}"
        viz.properties = _scrub_excel_artifacts(payload)
        viz.modified = now

    dashboard.modified = now
    core_db.info["audit_dw_id"] = system_dw_id

    from apps.app_dashboard.utils.audit_events import app_buffer_manual_audit
    app_buffer_manual_audit(
        core_db,
        action="UPDATE",
        module="dashboard",
        identifier=str(dashboard_id),
        new_value={
            "sections_updated": sections_updated,
            "html_fields_updated": html_updated,
            "visualization_fields_updated": viz_fields_updated,
            "visualizations_touched": len(props_by_viz),
            "languages": list(lang_col_by_code.keys()),
        },
        description=f"language import dashboard_id={dashboard_id} locales={list(lang_col_by_code.keys())}",
    )
    core_db.commit()

    return {
        "dashboard_id": dashboard_id,
        "sections_updated": sections_updated,
        "html_fields_updated": html_updated,
        "visualization_fields_updated": viz_fields_updated,
        "visualization_rows_skipped": viz_rows_skipped,
        "visualizations_touched": len(props_by_viz),
        "languages": list(lang_col_by_code.keys()),
    }
