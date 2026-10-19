"""
Dashboard language export — single-sheet XLSX for DFA_DASH v6.

Exports edit-dashboard section names, header/footer HTML, and visualization
translatable strings from the v6 nested `properties` shape across all DW
ut_language codes (default language first).
"""
from __future__ import annotations

import logging
import re
from datetime import datetime, timezone
from io import BytesIO
from typing import Any, Optional

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session

from apps.app_dashboard.constants import DASHBOARD_MODE_EDIT
from apps.app_dashboard.models.dashboard import (
    Dashboard,
    DashboardSection,
    DashVisualization,
)
from apps.app_dashboard.utils.dashboard_serialization import (
    parse_json_field,
    parse_json_or_raw_field,
)
from apps.datawarehouse.models.datawarehouse import Language
from apps.shared.exceptions import BadRequest, NotFound

logger = logging.getLogger(__name__)

_FILENAME_SAFE_RE = re.compile(r"[^A-Za-z0-9._-]+")

# Viz types that use FilterSettings (flat properties), not VisualizationSettings.
_FILTER_TYPES = frozenset({"select_list", "global"})
# Viz types that do not use the General title/axis block.
_NO_GENERAL_TYPES = frozenset({"html", "text", "image", "video", "tile", "select_list", "global"})


def _sanitize_filename_segment(name: Optional[str], fallback: str = "dashboard") -> str:
    raw = (name or "").strip() or fallback
    raw = raw.replace(" ", "-")
    safe = _FILENAME_SAFE_RE.sub("-", raw).strip("-")
    return safe or fallback


def _load_language_codes(dw_db: Session) -> tuple[list[str], Optional[str]]:
    """
    Return (ordered language codes, default language code).
    Default language first (language_default DESC), then language_code ASC.
    """
    rows = (
        dw_db.query(Language.language_code, Language.language_default)
        .order_by(Language.language_default.desc(), Language.language_code.asc())
        .all()
    )
    codes: list[str] = []
    default_code: Optional[str] = None
    for code, is_default in rows:
        if not code:
            continue
        c = str(code)
        codes.append(c)
        if is_default and default_code is None:
            default_code = c
    if codes and default_code is None:
        default_code = codes[0]
    return codes, default_code


def _i18n_row_values(
    field: Any,
    lang_codes: list[str],
    default_code: Optional[str],
) -> list[str]:
    """
    Spread a field across language columns:
    - dict → one cell per lang_codes order (missing → "")
    - string → value only under the default language column
    """
    if isinstance(field, dict):
        return ["" if field.get(code) is None else str(field.get(code)) for code in lang_codes]
    if field is None:
        return [""] * len(lang_codes)
    if isinstance(field, str):
        out = [""] * len(lang_codes)
        if default_code and default_code in lang_codes:
            out[lang_codes.index(default_code)] = field
        elif lang_codes:
            out[0] = field
        return out
    out = [""] * len(lang_codes)
    text = str(field)
    if default_code and default_code in lang_codes:
        out[lang_codes.index(default_code)] = text
    elif lang_codes:
        out[0] = text
    return out


def _append_lang_cells(ws, row: int, start_col: int, values: list[str]) -> None:
    for i, val in enumerate(values):
        ws.cell(row=row, column=start_col + i, value=val)


def _viz_type(viz: DashVisualization) -> str:
    return (viz.type or "").strip().lower()


def _append_if_present(
    rows: list[tuple[str, Optional[str], Any]],
    label: str,
    value: Any,
    *,
    key: Optional[str] = None,
) -> None:
    """Append a translatable field when the path exists (including empty string)."""
    if value is None:
        return
    rows.append((label, key, value))


def _collect_v6_viz_i18n_rows(
    props: dict[str, Any],
    viz_type: str,
) -> list[tuple[str, Optional[str], Any]]:
    """
    Collect (property_label, key_or_None, field) from v6 VisualizationSettings /
    FilterSettings stored in dash_visualization.properties.
    """
    rows: list[tuple[str, Optional[str], Any]] = []

    # ── Filter / select-list (flat FilterSettings) ────────────────────────
    if viz_type in _FILTER_TYPES:
        if "label" in props:
            _append_if_present(rows, "filterLabel", props.get("label"))
        return rows

    # ── Custom HTML ───────────────────────────────────────────────────────
    if viz_type == "html":
        custom_html = props.get("customHtml")
        if isinstance(custom_html, dict):
            if "html" in custom_html:
                _append_if_present(rows, "htmlText", custom_html.get("html"))
            if "noDataText" in custom_html:
                _append_if_present(rows, "htmlNoDataText", custom_html.get("noDataText"))
        return rows

    # ── Custom text ───────────────────────────────────────────────────────
    if viz_type == "text":
        custom_text = props.get("customText")
        if isinstance(custom_text, dict) and "text" in custom_text:
            _append_if_present(rows, "customText", custom_text.get("text"))
        return rows

    # image / video / tile — no author-entered prose in v6
    if viz_type in {"image", "video", "tile"}:
        return rows

    # ── Charts / map / table — General block ─────────────────────────────
    if viz_type not in _NO_GENERAL_TYPES:
        general = props.get("general")
        if isinstance(general, dict):
            chart_area = general.get("chartArea")
            if isinstance(chart_area, dict) and "noDataText" in chart_area:
                _append_if_present(rows, "noDataText", chart_area.get("noDataText"))

            title_block = general.get("title")
            if isinstance(title_block, dict):
                for entry_key, label in (
                    ("title", "title"),
                    ("subTitle", "subtitle"),
                    ("customLabels", "customLabel"),
                ):
                    entry = title_block.get(entry_key)
                    if isinstance(entry, dict) and "text" in entry:
                        _append_if_present(rows, label, entry.get("text"))

            axis_h = general.get("axisHorizontal")
            if isinstance(axis_h, dict) and "title" in axis_h:
                _append_if_present(rows, "xAxisTitle", axis_h.get("title"))

            axis_v = general.get("axisVertical")
            if isinstance(axis_v, dict) and "title" in axis_v:
                _append_if_present(rows, "yAxisTitle", axis_v.get("title"))

        # Series display-name overrides: properties.series.series[].name
        series_block = props.get("series")
        if isinstance(series_block, dict):
            series_list = series_block.get("series")
            if isinstance(series_list, list):
                for idx, item in enumerate(series_list):
                    if isinstance(item, dict) and "name" in item:
                        key = str(item.get("id", idx))
                        _append_if_present(
                            rows,
                            f"series.{key}",
                            item.get("name"),
                            key="name",
                        )

        # Category display-name overrides: properties.categories.categories[].name
        categories_block = props.get("categories")
        if isinstance(categories_block, dict):
            categories_list = categories_block.get("categories")
            if isinstance(categories_list, list):
                for idx, item in enumerate(categories_list):
                    if isinstance(item, dict) and "name" in item:
                        key = str(item.get("id", item.get("index", idx)))
                        _append_if_present(
                            rows,
                            f"categories.{key}",
                            item.get("name"),
                            key="name",
                        )

    return rows


def build_dashboard_languages_export_xlsx(
    core_db: Session,
    dw_db: Session,
    *,
    dashboard_id: int,
    system_dw_id: int,
) -> tuple[BytesIO, str]:
    """
    Build dashboard translation workbook for v6 properties.

    Raises NotFound / BadRequest on validation failures.
    """
    dashboard = core_db.query(Dashboard).filter(Dashboard.id == dashboard_id).first()
    if not dashboard:
        raise NotFound(
            f"Dashboard with id {dashboard_id} not found",
            error_code="DASHBOARD_NOT_FOUND",
        )
    if dashboard.mode != DASHBOARD_MODE_EDIT:
        raise BadRequest(
            "Only edit-mode dashboards can be exported",
            error_code="DASHBOARD_NOT_EDIT_MODE",
            details={"dashboard_id": dashboard_id, "mode": dashboard.mode},
        )
    if dashboard.system_dw_id != system_dw_id:
        raise BadRequest(
            "Dashboard is not linked to the requested data warehouse",
            error_code="DASHBOARD_DW_MISMATCH",
            details={
                "dashboard_id": dashboard_id,
                "dashboard_system_dw_id": dashboard.system_dw_id,
                "path_system_dw_id": system_dw_id,
            },
        )

    lang_codes, default_code = _load_language_codes(dw_db)
    if not lang_codes:
        raise BadRequest(
            "No languages configured for this data warehouse",
            error_code="DW_LANGUAGES_EMPTY",
        )

    sections = (
        core_db.query(DashboardSection)
        .filter(DashboardSection.dashboard_id == dashboard_id)
        .order_by(DashboardSection.order.asc(), DashboardSection.id.asc())
        .all()
    )
    visualizations = (
        core_db.query(DashVisualization)
        .filter(
            DashVisualization.dashboard_id == dashboard_id,
            DashVisualization.is_deleted == 0,
            DashVisualization.is_used_in_layout == 1,
        )
        .order_by(DashVisualization.id.asc())
        .all()
    )

    wb = Workbook()
    ws = wb.active
    ws.title = "Worksheet"

    # Header row
    ws.cell(row=1, column=1, value="Property Name")
    ws.cell(row=1, column=2, value="Section Name")
    for i, code in enumerate(lang_codes):
        ws.cell(row=1, column=5 + i, value=code)

    row = 2

    # ── Sections ──────────────────────────────────────────────────────────
    for section in sections:
        section_name = parse_json_or_raw_field(section.section_name)
        if section_name is None:
            continue
        ws.cell(row=row, column=1, value="Section Names")
        _append_lang_cells(ws, row, 5, _i18n_row_values(section_name, lang_codes, default_code))
        row += 1

    # ── Dashboard HTML ────────────────────────────────────────────────────
    for label, raw in (
        ("Header HTML", dashboard.header_html),
        ("Mobile Header HTML", dashboard.mobile_header_html),
        ("Footer HTML", dashboard.footer_html),
    ):
        value = parse_json_or_raw_field(raw)
        ws.cell(row=row, column=1, value=label)
        if value is not None:
            _append_lang_cells(ws, row, 5, _i18n_row_values(value, lang_codes, default_code))
        row += 1

    # Blank separator
    row += 1

    # Viz header
    ws.cell(row=row, column=1, value="Visualization IDs")
    ws.cell(row=row, column=2, value="Visualization Name")
    ws.cell(row=row, column=3, value="Properties")
    ws.cell(row=row, column=4, value="Key")
    row += 1

    for viz in visualizations:
        props = parse_json_field(viz.properties) or {}
        if not isinstance(props, dict):
            props = {}

        viz_rows = _collect_v6_viz_i18n_rows(props, _viz_type(viz))
        if not viz_rows:
            continue

        viz_name = viz.visualization_name or ""
        for prop_label, key, field in viz_rows:
            ws.cell(row=row, column=1, value=viz.id)
            ws.cell(row=row, column=2, value=viz_name)
            ws.cell(row=row, column=3, value=prop_label)
            if key is not None:
                ws.cell(row=row, column=4, value=key)
            _append_lang_cells(ws, row, 5, _i18n_row_values(field, lang_codes, default_code))
            row += 1

        # Blank row between visualizations
        row += 1

    # Styling: bold header row + italic/bold on A–D
    bold = Font(bold=True)
    bold_italic = Font(bold=True, italic=True)
    for col in range(1, 5 + len(lang_codes)):
        ws.cell(row=1, column=col).font = bold
    for r in range(1, row):
        for col in range(1, 5):
            cell = ws.cell(row=r, column=col)
            if cell.value is not None:
                cell.font = bold_italic
    for col in range(1, 5 + len(lang_codes)):
        ws.cell(row=1, column=col).font = bold

    # Auto-size columns to max content (capped for readability)
    max_col = 4 + len(lang_codes)
    max_lens = [0] * max_col
    for r in range(1, ws.max_row + 1):
        for c in range(1, max_col + 1):
            val = ws.cell(row=r, column=c).value
            if val is None:
                continue
            text = str(val)
            length = max((len(line) for line in text.splitlines()), default=0)
            if length > max_lens[c - 1]:
                max_lens[c - 1] = length

    padding = 6
    min_width = 12
    max_width = 80
    for i, max_len in enumerate(max_lens, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(
            max(max_len + padding, min_width),
            max_width,
        )

    ts = datetime.now(timezone.utc).strftime("%Y-%m-%d-%H-%M-%S")
    filename = f"{_sanitize_filename_segment(dashboard.dashboard_name)}-{ts}.xlsx"

    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output, filename
